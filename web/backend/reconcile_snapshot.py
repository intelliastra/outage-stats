"""Compare the active database dataset with a historical result snapshot."""

from __future__ import annotations

import argparse
import hashlib
import json
from collections import Counter
from datetime import date, datetime
from pathlib import Path
from typing import Any

import openpyxl

import database


CALCULATED_COLUMNS = {"用户停电总次数", "是否统计", "不统计原因", "频繁停电类型", "停电预警类型"}


def normalized(value: Any) -> Any:
    if value is None:
        return None
    if isinstance(value, (date, datetime)):
        return value.isoformat()
    if isinstance(value, float) and value.is_integer():
        return int(value)
    text = str(value).strip()
    return None if text.lower() in {"", "nan", "nat", "none"} else text


def row_hash(payload: dict[str, Any], columns: list[str]) -> str:
    values = [normalized(payload.get(column)) for column in columns]
    canonical = json.dumps(values, ensure_ascii=False, separators=(",", ":"))
    return hashlib.sha256(canonical.encode("utf-8")).hexdigest()


def snapshot_rows(path: Path):
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    worksheet = workbook["用户停电总次数统计表"] if "用户停电总次数统计表" in workbook.sheetnames else workbook.worksheets[0]
    rows = worksheet.iter_rows(values_only=True)
    headers = [str(value).strip() if value is not None else "" for value in next(rows)]
    try:
        for values in rows:
            if not any(value is not None for value in values):
                continue
            yield {header: value for header, value in zip(headers, values) if header}
    finally:
        workbook.close()


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("snapshot", type=Path)
    parser.add_argument("--report", type=Path, required=True)
    args = parser.parse_args()

    source_rows = list(snapshot_rows(args.snapshot))
    if not source_rows:
        raise SystemExit("snapshot has no rows")
    snapshot_columns = [column for column in source_rows[0] if column not in CALCULATED_COLUMNS]

    with database.connection() as conn:
        db_rows = conn.execute("SELECT raw_payload FROM current_outage_record ORDER BY id").fetchall()
    active_rows = [row["raw_payload"] for row in db_rows]
    common_columns = [column for column in snapshot_columns if any(column in row for row in active_rows[:1000])]
    snapshot_hashes = Counter(row_hash(row, common_columns) for row in source_rows)
    active_hashes = Counter(row_hash(row, common_columns) for row in active_rows)
    missing = snapshot_hashes - active_hashes
    extra = active_hashes - snapshot_hashes
    report = {
        "snapshot": str(args.snapshot.resolve()),
        "snapshot_rows": len(source_rows),
        "database_rows": len(active_rows),
        "columns_compared": common_columns,
        "missing_rows": sum(missing.values()),
        "extra_rows": sum(extra.values()),
        "match": not missing and not extra,
        "missing_hash_samples": list(missing)[:100],
        "extra_hash_samples": list(extra)[:100],
    }
    args.report.parent.mkdir(parents=True, exist_ok=True)
    args.report.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(report, ensure_ascii=False))
    return 0 if report["match"] else 2


if __name__ == "__main__":
    raise SystemExit(main())

