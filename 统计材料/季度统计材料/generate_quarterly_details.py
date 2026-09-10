#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""按季度输出停电统计明细清单（支持剔除口径 / 全量口径，含 NEA 标记）。"""

from __future__ import annotations

import argparse
import sys
from dataclasses import dataclass
from datetime import date
from pathlib import Path

import openpyxl
import pandas as pd
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

SCRIPT_DIR = Path(__file__).resolve().parent
sys.path.insert(0, str(SCRIPT_DIR))

import generate_quarterly_stats as gq

OUTPUT_PATH_EXCLUDE = SCRIPT_DIR / "2025年季度停电统计明细.xlsx"
OUTPUT_PATH_FULL = SCRIPT_DIR / "2025年季度停电统计明细_全量.xlsx"
SUMMARY_PATH = SCRIPT_DIR / "2025年季度停电统计数据.xlsx"

USER_COLUMNS = [
    "用户编码",
    "中电联用户名称",
    "用户性质",
    "所属馈线编码",
    "所属馈线名称",
    "所属供电所",
    "所属区局",
    "所属地市",
    "停电总次数",
    "故障停电次数",
    "预安排停电次数",
    "频繁停电类型",
    "停电预警类型",
]

LINE_COLUMNS = [
    "所属馈线编码",
    "所属馈线名称",
    "所属供电所",
    "所属区局",
    "所属地市",
    "停电总次数",
    "故障停电次数",
    "预安排停电次数",
    "频繁停电类型",
    "停电预警类型",
    "是否纳入国家能源局台账",
]

WARNING_LINE_COLUMNS = LINE_COLUMNS + ["预警时间", "是否频繁停"]

TEXT_COLUMNS = {"用户编码", "所属馈线编码"}
NEA_BANNER = "【NEA·已纳入国家能源局管控线路】统计区间：{start} ~ {end}"
NEA_COL_HEADER = "【NEA】是否纳入国家能源局台账"


@dataclass(frozen=True)
class QuarterDef:
    key: str
    label: str
    start: date
    end: date
    year: int


@dataclass(frozen=True)
class DetailSheetDef:
    sheet_name: str
    quarter_key: str
    list_key: str
    columns: list[str]
    is_nea: bool = False
    summary_cols: str = ""


QUARTERS: list[QuarterDef] = [
    QuarterDef("2025Q1", "2025年第一季度", date(2025, 1, 1), date(2025, 3, 31), 2025),
    QuarterDef("2025Q2", "2025年第二季度（累计值）", date(2025, 1, 1), date(2025, 6, 30), 2025),
    QuarterDef("2025Q3", "2025年第三季度（累计值）", date(2025, 1, 1), date(2025, 9, 30), 2025),
    QuarterDef("2025Q4", "2025年第四季度（累计值）", date(2025, 1, 1), date(2025, 12, 31), 2025),
    QuarterDef("2026Q1", "2026年第一季度", date(2026, 1, 1), date(2026, 3, 31), 2026),
]


def build_detail_sheet_defs() -> list[DetailSheetDef]:
    sheets: list[DetailSheetDef] = []
    for q in QUARTERS:
        if q.year == 2025:
            sheets.extend(
                [
                    DetailSheetDef(
                        f"{q.key}-频繁停电用户",
                        q.key,
                        "频繁停电用户清单",
                        USER_COLUMNS,
                        summary_cols="B-E",
                    ),
                    DetailSheetDef(
                        f"{q.key}-频繁停电线路",
                        q.key,
                        "频繁停电线路清单",
                        LINE_COLUMNS,
                        summary_cols="F-I",
                    ),
                    DetailSheetDef(
                        f"{q.key}-NEA频繁停电线路",
                        q.key,
                        "频繁停电线路清单",
                        LINE_COLUMNS,
                        is_nea=True,
                        summary_cols="J-M",
                    ),
                ]
            )
        else:
            sheets.extend(
                [
                    DetailSheetDef(
                        f"{q.key}-预警用户",
                        q.key,
                        "停电预警用户清单",
                        USER_COLUMNS,
                        summary_cols="AY-BB",
                    ),
                    DetailSheetDef(
                        f"{q.key}-预警线路",
                        q.key,
                        "停电预警线路清单",
                        WARNING_LINE_COLUMNS,
                        summary_cols="BC-BF",
                    ),
                    DetailSheetDef(
                        f"{q.key}-频繁停电用户",
                        q.key,
                        "频繁停电用户清单",
                        USER_COLUMNS,
                        summary_cols="BG-BJ",
                    ),
                    DetailSheetDef(
                        f"{q.key}-频繁停电线路",
                        q.key,
                        "频繁停电线路清单",
                        LINE_COLUMNS,
                        summary_cols="BK-BN",
                    ),
                    DetailSheetDef(
                        f"{q.key}-NEA频繁停电线路",
                        q.key,
                        "频繁停电线路清单",
                        LINE_COLUMNS,
                        is_nea=True,
                        summary_cols="BO-BR",
                    ),
                ]
            )
    return sheets


DETAIL_SHEETS = build_detail_sheet_defs()


def quarter_by_key(key: str) -> QuarterDef:
    for q in QUARTERS:
        if q.key == key:
            return q
    raise KeyError(key)


def extract_list_df(
    tables: dict[str, pd.DataFrame],
    list_key: str,
    columns: list[str],
    *,
    is_nea: bool,
) -> pd.DataFrame:
    df = tables.get(list_key, pd.DataFrame()).copy()
    if df.empty:
        return pd.DataFrame(columns=columns)

    if is_nea:
        nea_col = "是否纳入国家能源局台账"
        if nea_col in df.columns:
            df = df[df[nea_col].astype(str).str.strip().eq("是")].copy()
        else:
            df = df.iloc[0:0].copy()

    for col in columns:
        if col not in df.columns:
            df[col] = ""
    df = df[columns].copy()

    if is_nea and "是否纳入国家能源局台账" in df.columns:
        df["是否纳入国家能源局台账"] = "是"

    for col in TEXT_COLUMNS & set(columns):
        if col in df.columns:
            df[col] = (
                df[col]
                .astype(str)
                .str.strip()
                .str.replace(".0", "", regex=False)
            )

    return df.reset_index(drop=True)


def header_labels(columns: list[str], *, is_nea: bool) -> list[str]:
    labels = []
    for col in columns:
        if is_nea and col == "是否纳入国家能源局台账":
            labels.append(NEA_COL_HEADER)
        else:
            labels.append(col)
    return labels


def write_detail_sheet(
    ws,
    df: pd.DataFrame,
    columns: list[str],
    *,
    is_nea: bool,
    period_start: date,
    period_end: date,
) -> None:
    labels = header_labels(columns, is_nea=is_nea)
    header_row = 2 if is_nea else 1
    data_start = header_row + 1

    if is_nea:
        banner = NEA_BANNER.format(start=period_start.isoformat(), end=period_end.isoformat())
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=max(len(columns), 1))
        cell = ws.cell(1, 1, banner)
        cell.font = Font(name="微软雅黑", bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    center_font = Font(name="微软雅黑")
    center_align = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for col_idx, label in enumerate(labels, start=1):
        c = ws.cell(header_row, col_idx, label)
        c.font = Font(name="微软雅黑", bold=True)
        c.alignment = center_align

    for row_offset, row in enumerate(df.itertuples(index=False), start=data_start):
        for col_idx, value in enumerate(row, start=1):
            cell = ws.cell(row_offset, col_idx, value if pd.notna(value) and str(value).strip() else None)
            cell.font = center_font
            cell.alignment = center_align
            col_name = columns[col_idx - 1]
            if col_name in TEXT_COLUMNS:
                cell.number_format = "@"

    for col_idx, col_name in enumerate(columns, start=1):
        letter = get_column_letter(col_idx)
        width = max(len(col_name) + 2, 12)
        if col_name in {"频繁停电类型", "停电预警类型", "所属馈线名称", "中电联用户名称"}:
            width = min(width + 8, 36)
        ws.column_dimensions[letter].width = width

    if is_nea:
        ws.row_dimensions[1].height = 28


def build_meta_sheet(wb, *, exclude_major: bool, major_dates: set[date]) -> None:
    ws = wb.create_sheet("明细说明")
    caliber = "剔除口径（剔除表11重大事件日）" if exclude_major else "全量口径（不剔除表11重大事件日）"
    lines = [
        "明细说明",
        "",
        f"1. 统计口径：{caliber}",
        "2. 数据来源：2025年停电用户数据.xlsx、2026.1.1-6.16停电用户_处理结果.xlsx",
    ]
    if exclude_major:
        lines.extend(
            [
                "3. 重大事件日剔除：表11重大事件日统计表（日期见下）",
                f"   剔除日期：{', '.join(sorted(d.isoformat() for d in major_dates))}",
            ]
        )
    else:
        lines.append("3. 重大事件日：不剔除（全量统计，表11日期均保留计入）")
        if major_dates:
            lines.append(
                f"   表11列示日期（本全量版未剔除）：{', '.join(sorted(d.isoformat() for d in major_dates))}"
            )

    lines.extend(
        [
            "",
            "4. 各季度统计区间：",
            "   2025Q1：2025-01-01 ~ 2025-03-31",
            "   2025Q2累计：2025-01-01 ~ 2025-06-30",
            "   2025Q3累计：2025-01-01 ~ 2025-09-30",
            "   2025Q4累计：2025-01-01 ~ 2025-12-31",
            "   2026Q1：2026-01-01 ~ 2026-03-31",
            "",
            "5. Sheet 与汇总列对照（剔除版汇总表对应剔除明细；全量明细需与全量汇总对照）：",
        ]
    )
    for spec in DETAIL_SHEETS:
        nea_tag = "（NEA·国家能源局管控线路）" if spec.is_nea else ""
        lines.append(f"   {spec.sheet_name}{nea_tag} → 汇总列 {spec.summary_cols}")

    lines.extend(
        [
            "",
            "6. NEA 明细说明：",
            "   - Sheet 名含 NEA，第1行横幅标注【NEA·已纳入国家能源局管控线路】",
            "   - 台账列标题为【NEA】是否纳入国家能源局台账，数据均为「是」",
            "   - NEA 线路为全量频繁停电线路的子集，对应汇总 J-M / BO-BR 列",
            "",
            "7. 合计列（E/I/M 等）= 三个子指标之和，子指标可重叠，不等于明细去重行数",
        ]
    )

    font = Font(name="微软雅黑")
    for row_idx, text in enumerate(lines, start=1):
        ws.cell(row_idx, 1, text).font = font
    ws.column_dimensions["A"].width = 100


def count_by_keyword(df: pd.DataFrame, type_col: str, keyword: str, city_col: str = "所属地市") -> dict[str, int]:
    if df.empty or type_col not in df.columns:
        return {}
    sub = df[df[type_col].fillna("").astype(str).str.contains(keyword, na=False, regex=False)]
    if sub.empty:
        return {}
    return sub.groupby(city_col).size().to_dict()


def verify_2025q1(mod, tables: dict[str, pd.DataFrame], *, compare_summary: bool) -> None:
    print("\n=== 2025Q1 抽样校验 ===", flush=True)
    dicts = mod.compute_2025_stats_dicts(tables)
    fu = tables["频繁停电用户清单"]
    fl = tables["频繁停电线路清单"]
    fl_nea = extract_list_df(tables, "频繁停电线路清单", LINE_COLUMNS, is_nea=True)

    checks = [
        ("用户-超5次", dicts[0], count_by_keyword(fu, "频繁停电类型", "一年内停电次数超过5次"), 2),
        ("用户-60天超3次", dicts[1], count_by_keyword(fu, "频繁停电类型", "连续60天停电次数超过3次"), 3),
        ("线路-超5次", dicts[4], count_by_keyword(fl, "频繁停电类型", "一年内停电次数超过5次"), 6),
        ("NEA线路-超5次", dicts[8], count_by_keyword(fl_nea, "频繁停电类型", "一年内停电次数超过5次"), 10),
    ]

    summary_vals: dict[int, int | None] = {}
    if compare_summary and SUMMARY_PATH.exists():
        summary_wb = openpyxl.load_workbook(SUMMARY_PATH, read_only=True)
        summary_ws = summary_wb["停电统计数据"]
        for _, _, _, col in checks:
            summary_vals[col] = summary_ws.cell(4, col).value
        summary_wb.close()

    for name, dict_from_stats, dict_from_detail, sum_col in checks:
        detail_total = sum(dict_from_detail.values())
        stats_total = sum(dict_from_stats.values())
        summary_val = summary_vals.get(sum_col) if compare_summary else "—"
        ok = stats_total == detail_total
        print(
            f"  {name}: 明细={detail_total}, 统计={stats_total}, 汇总表={summary_val} "
            f"{'OK' if ok else 'MISMATCH'}",
            flush=True,
        )

    print(
        f"  NEA子集: 全量线路 {len(fl)} 条, NEA线路 {len(fl_nea)} 条 "
        f"({'OK' if len(fl_nea) <= len(fl) else 'MISMATCH'})",
        flush=True,
    )


def generate_details(*, exclude_major: bool, output_path: Path) -> None:
    mod = gq.load_process_module()
    major_dates = gq.read_major_event_dates_from_table11(gq.MAJOR_EVENT_PATH)
    major_2025 = {d for d in major_dates if d.year == 2025} if exclude_major else set()
    major_2026 = {d for d in major_dates if d.year == 2026} if exclude_major else set()
    nea_gis_ids = mod.read_nea_ledger(str(gq.NEA_LEDGER_PATH)) if gq.NEA_LEDGER_PATH.exists() else set()

    caliber_label = "剔除口径" if exclude_major else "全量口径"
    print(f"\n========== 生成明细：{caliber_label} ==========", flush=True)

    print("预加载 2025 年原始数据...", flush=True)
    raw_2025 = mod.normalize_2025_columns(mod.read_2025_raw(str(gq.RAW_2025_PATH)))
    print(f"  2025 年有效记录：{len(raw_2025)} 行", flush=True)

    print("预加载 2026 年原始数据...", flush=True)
    raw_2026 = gq.read_2026_raw(gq.RAW_2026_PATH, mod)
    print(f"  2026 年有效记录：{len(raw_2026)} 行", flush=True)

    quarter_tables: dict[str, dict[str, pd.DataFrame]] = {}
    for q in QUARTERS:
        print(f"计算 {q.label}（{q.start} ~ {q.end}）...", flush=True)
        raw_base = raw_2025 if q.year == 2025 else raw_2026
        period_major = major_2025 if q.year == 2025 else major_2026
        quarter_tables[q.key] = gq.build_period_tables(
            mod,
            raw_base=raw_base,
            start=q.start,
            end=q.end,
            period_major=period_major,
            nea_gis_ids=nea_gis_ids,
        )

    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    for spec in DETAIL_SHEETS:
        q = quarter_by_key(spec.quarter_key)
        tables = quarter_tables[spec.quarter_key]
        df = extract_list_df(tables, spec.list_key, spec.columns, is_nea=spec.is_nea)
        ws = wb.create_sheet(spec.sheet_name)
        write_detail_sheet(
            ws,
            df,
            spec.columns,
            is_nea=spec.is_nea,
            period_start=q.start,
            period_end=q.end,
        )
        nea_note = " [NEA]" if spec.is_nea else ""
        print(f"  写入 {spec.sheet_name}{nea_note}: {len(df)} 行", flush=True)

    build_meta_sheet(wb, exclude_major=exclude_major, major_dates=major_dates)
    wb.save(output_path)
    print(f"已输出：{output_path}", flush=True)

    verify_2025q1(mod, quarter_tables["2025Q1"], compare_summary=exclude_major)


def main() -> None:
    parser = argparse.ArgumentParser(description="生成季度停电统计明细 Excel")
    parser.add_argument(
        "--mode",
        choices=["exclude", "full", "both"],
        default="exclude",
        help="exclude=剔除表11重大事件日; full=全量; both=两种都生成",
    )
    args = parser.parse_args()

    if args.mode in ("exclude", "both"):
        generate_details(exclude_major=True, output_path=OUTPUT_PATH_EXCLUDE)
    if args.mode in ("full", "both"):
        generate_details(exclude_major=False, output_path=OUTPUT_PATH_FULL)


if __name__ == "__main__":
    main()
