#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""按季度累计口径生成停电统计数据（剔除表11重大事件日）。"""

from __future__ import annotations

import importlib.util
import sys
from datetime import date
from pathlib import Path

import openpyxl
import pandas as pd
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

SCRIPT_DIR = Path(__file__).resolve().parent
PROJECT_SCRIPT = SCRIPT_DIR.parent / "process_outage_tables_v1.0.3.py"

RAW_2025_PATH = SCRIPT_DIR / "2025年停电用户数据.xlsx"
RAW_2026_PATH = SCRIPT_DIR / "2026.1.1-6.16停电用户_处理结果.xlsx"
MAJOR_EVENT_PATH = SCRIPT_DIR / "表11重大事件日统计表_20260618085330546.xlsx"
NEA_LEDGER_PATH = SCRIPT_DIR.parent / "国家能源局台账.xlsx"
STATS_TEMPLATE_PATH = SCRIPT_DIR.parent / "统计表（模板）.xlsx"
OUTPUT_PATH = SCRIPT_DIR / "季度停电统计数据.xlsx"

QUARTER_HEADERS = [
    "2025年第一季度",
    "2025年第二季度（累计值）",
    "2025年第三季度（累计值）",
    None,
    "2025年第四季度（累计值）",
    "2026年第一季度",
]

PERIODS = [
    (date(2025, 1, 1), date(2025, 3, 31), 2025, "2025"),
    (date(2025, 1, 1), date(2025, 6, 30), 2025, "2025"),
    (date(2025, 1, 1), date(2025, 9, 30), 2025, "2025"),
    None,
    (date(2025, 1, 1), date(2025, 12, 31), 2025, "2025"),
    (date(2026, 1, 1), date(2026, 3, 31), 2026, "2026"),
]

GROUP_HEADERS_2025 = [
    "2025年频繁停电用户",
    "2025年频繁停电线路",
    "2025年频繁停电线路\n（已纳入国家能源局管控线路）",
]

GROUP_HEADERS_2026 = [
    "2026年频繁停电预警用户",
    "2026年频繁停电预警线路",
    "2026年频繁停电用户",
    "2026年频繁停电线路",
    "2026年频繁停电线路\n（已纳入国家能源局管控线路）",
]

SUB_HEADERS_WARNING = [
    "一年内停电4-5次",
    "近50天停电3次",
    "近30天停电2次",
    "合计",
]

SUB_HEADERS_FREQ = [
    "停电总次数超过5次",
    "连续60天停电次数超过3次",
    "一年内预安排停电次数超过3次",
    "合计",
]


def load_process_module():
    spec = importlib.util.spec_from_file_location("process_outage_tables", PROJECT_SCRIPT)
    module = importlib.util.module_from_spec(spec)
    sys.modules["process_outage_tables"] = module
    spec.loader.exec_module(module)
    return module


def read_major_event_dates_from_table11(path: Path) -> set[date]:
    df = pd.read_excel(path)
    col = next((c for c in df.columns if "重大事件日期" in str(c)), None)
    if col is None:
        raise KeyError("表11中未找到「重大事件日期」列")
    dates: set[date] = set()
    for value in df[col]:
        parsed = pd.to_datetime(value, errors="coerce")
        if pd.notna(parsed):
            dates.add(parsed.date())
    return dates


def read_2026_raw(path: Path, mod) -> pd.DataFrame:
    df = pd.read_excel(path, sheet_name="用户停电总次数统计表", dtype=str)
    drop_cols = [
        c
        for c in mod.RAW_ADDED_COLUMNS + ["_user_feeder_key", "_是否统计_bool", "_不统计原因"]
        if c in df.columns
    ]
    df = df.drop(columns=drop_cols)
    df.columns = [str(c).strip() for c in df.columns]
    return df.dropna(how="all").reset_index(drop=True)


def build_period_tables(
    mod,
    *,
    raw_base: pd.DataFrame,
    start: date,
    end: date,
    period_major: set[date],
    nea_gis_ids: set[str],
) -> dict[str, pd.DataFrame]:
    raw = mod.filter_raw_by_outage_date(raw_base, start, end)
    tables, _ = mod.build_tables(
        raw,
        end,
        period_major,
        previous_repeated_lines=pd.DataFrame(columns=mod.REPEATED_LINE_COLUMNS),
        nea_gis_ids=nea_gis_ids,
    )
    return tables


def get_template_cities(template_path: Path) -> list[str]:
    wb = openpyxl.load_workbook(template_path, read_only=True)
    ws = wb["2025年统计表"] if "2025年统计表" in wb.sheetnames else wb.active
    cities = []
    for row in ws.iter_rows(min_row=3, max_col=1, values_only=True):
        value = row[0]
        if value is not None and str(value).strip():
            cities.append(str(value).strip())
    wb.close()
    return cities


def company_total(dicts: list[dict[str, int]], cities: list[str], mod) -> list[int]:
    totals = []
    for d in dicts:
        total = 0
        for city_name in cities:
            if "公司" in city_name:
                continue
            matched = mod._match_city(city_name, list(d.keys()))
            if matched:
                total += int(d.get(matched, 0))
        totals.append(total)
    return totals


def fill_city_values(
    ws,
    row_idx: int,
    col_start: int,
    dicts: list[dict[str, int]],
    city_name: str,
    template_cities: list[str],
    all_cities: list[str],
    mod,
) -> None:
    if "公司" in city_name:
        values = company_total(dicts, template_cities, mod)
    else:
        matched = mod._match_city(city_name, all_cities)
        values = [int(d.get(matched, 0)) if matched else 0 for d in dicts]

    for offset, val in enumerate(values):
        ws.cell(row_idx, col_start + offset).value = val


def period_layout(fmt: str) -> tuple[int, list[str], list[str]]:
    if fmt == "2026":
        groups = GROUP_HEADERS_2026
        subs = SUB_HEADERS_WARNING[:4] + SUB_HEADERS_WARNING[:4] + SUB_HEADERS_FREQ[:4] * 3
        return 20, groups, subs
    groups = GROUP_HEADERS_2025
    subs = SUB_HEADERS_FREQ[:4] * 3
    return 12, groups, subs


def build_workbook(
    period_dicts: list[tuple[str | None, list[dict[str, int]] | None, str | None]],
    cities: list[str],
    mod,
    major_dates: set[date],
) -> openpyxl.Workbook:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "停电统计数据"

    col = 2
    period_col_ranges: list[tuple[int, int] | None] = []

    ws.cell(1, 1, "单位")
    ws.cell(2, 1, None)
    ws.cell(3, 1, None)

    for header, period_data in zip(QUARTER_HEADERS, period_dicts):
        if header is None:
            period_col_ranges.append(None)
            col += 1
            continue

        _, period_dict, fmt = period_data
        metrics_per_period, groups, subs = period_layout(fmt or "2025")
        start_col = col
        end_col = col + metrics_per_period - 1
        period_col_ranges.append((start_col, end_col))

        ws.merge_cells(start_row=1, start_column=start_col, end_row=1, end_column=end_col)
        ws.cell(1, start_col, header)

        sub_per_group = 4
        for group_idx, group_name in enumerate(groups):
            group_start = col + group_idx * sub_per_group
            group_end = group_start + sub_per_group - 1
            ws.merge_cells(start_row=2, start_column=group_start, end_row=2, end_column=group_end)
            ws.cell(2, group_start, group_name)
            group_subs = SUB_HEADERS_WARNING if fmt == "2026" and group_idx < 2 else SUB_HEADERS_FREQ
            for sub_idx, sub_name in enumerate(group_subs):
                ws.cell(3, group_start + sub_idx, sub_name)

        col = end_col + 1

    for row_offset, city_name in enumerate(cities):
        row_idx = 4 + row_offset
        ws.cell(row_idx, 1, city_name)
        for period_range, period_data in zip(period_col_ranges, period_dicts):
            if period_range is None or period_data[1] is None:
                continue
            start_col, _ = period_range
            period_dict = period_data[1]
            all_cities = list({k for d in period_dict for k in d})
            fill_city_values(ws, row_idx, start_col, period_dict, city_name, cities, all_cities, mod)

    center_font = Font(name="微软雅黑")
    center_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for row in ws.iter_rows():
        for cell in row:
            cell.font = center_font
            cell.alignment = center_align

    ws.column_dimensions["A"].width = 12
    for c in range(2, col):
        ws.column_dimensions[get_column_letter(c)].width = 14

    meta = wb.create_sheet("统计说明")
    meta["A1"] = "统计口径说明"
    meta["A2"] = "1. 数据来源：2025年停电用户数据.xlsx、2026.1.1-6.16停电用户_处理结果.xlsx"
    meta["A3"] = "2. 重大事件日剔除：表11重大事件日统计表_20260618085330546.xlsx"
    meta["A4"] = f"3. 剔除日期：{', '.join(sorted(d.isoformat() for d in major_dates))}"
    meta["A5"] = "4. 2025年第二至四季度为年初至该季度末的累计值"
    meta["A6"] = "5. 2026年第一季度统计区间：2026-01-01至2026-03-31，采用2026年统计表口径（含预警用户/线路）"
    meta["A7"] = "6. 合计列=三个子项之和（子项可重叠）；其余规则与 process_outage_tables_v1.0.3 一致"

    return wb


def main() -> None:
    mod = load_process_module()
    major_dates = read_major_event_dates_from_table11(MAJOR_EVENT_PATH)
    major_2025 = {d for d in major_dates if d.year == 2025}
    major_2026 = {d for d in major_dates if d.year == 2026}
    nea_gis_ids = mod.read_nea_ledger(str(NEA_LEDGER_PATH)) if NEA_LEDGER_PATH.exists() else set()
    cities = get_template_cities(STATS_TEMPLATE_PATH)

    print("预加载 2025 年原始数据...")
    raw_2025 = mod.normalize_2025_columns(mod.read_2025_raw(str(RAW_2025_PATH)))
    print(f"  2025 年有效记录：{len(raw_2025)} 行")

    print("预加载 2026 年原始数据...")
    raw_2026 = read_2026_raw(RAW_2026_PATH, mod)
    print(f"  2026 年有效记录：{len(raw_2026)} 行")

    period_dicts: list[tuple[str | None, list[dict[str, int]] | None, str | None]] = []
    for header, period in zip(QUARTER_HEADERS, PERIODS):
        if header is None or period is None:
            period_dicts.append((header, None, None))
            continue

        start, end, year, fmt = period
        print(f"计算 {header}（{start} ~ {end}）...", flush=True)
        raw_base = raw_2025 if year == 2025 else raw_2026
        period_major = major_2025 if year == 2025 else major_2026
        tables = build_period_tables(
            mod,
            raw_base=raw_base,
            start=start,
            end=end,
            period_major=period_major,
            nea_gis_ids=nea_gis_ids,
        )
        if fmt == "2026":
            dicts = mod.compute_2026_stats_dicts(tables)
            company_vals = company_total(dicts, cities, mod)
            print(
                f"  公司合计（预警用户/预警线路/频繁用户/频繁线路）："
                f"{company_vals[3]} / {company_vals[7]} / {company_vals[11]} / {company_vals[15]}",
                flush=True,
            )
        else:
            dicts = mod.compute_2025_stats_dicts(tables)
            company_vals = company_total(dicts, cities, mod)
            print(
                f"  公司合计（用户/线路/NEA 合计列）："
                f"{company_vals[3]} / {company_vals[7]} / {company_vals[11]}",
                flush=True,
            )
        period_dicts.append((header, dicts, fmt))

    wb = build_workbook(period_dicts, cities, mod, major_dates)
    wb.save(OUTPUT_PATH)
    print(f"\n已输出：{OUTPUT_PATH}")


if __name__ == "__main__":
    main()
