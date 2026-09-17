#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""Generate outage reliability rule tables from an exported user outage workbook.

Version: 1.0.8 — Newdata closed date-window replaces overlapping output rows;
days missing inside the window stay empty (no backfill). History before/after
the window is kept; stats use the merged year-to-date range.
"""

from __future__ import annotations

# ============================================================
# ▶▶▶  路径配置区域（根据实际情况修改，其余留空自动命名）  ◀◀◀
# ============================================================
import os as _os
_SCRIPT_DIR = _os.path.dirname(_os.path.abspath(__file__))

# --- 输入路径（日常统计口径） ---
NEWDATA_FOLDER       = _os.path.join(_SCRIPT_DIR, "input", "Newdata", "daily")
NEA_LEDGER_PATH      = _os.path.join(_SCRIPT_DIR, "国家能源局台账.xlsx")
STATS_TEMPLATE_PATH  = _os.path.join(_SCRIPT_DIR, "统计表（模板）.xlsx")
RAW_2025_DATA_PATH   = _os.path.join(_SCRIPT_DIR, "2025年停电用户数据.xlsx")
ANNUAL_2025_PATH     = _os.path.join(
    _SCRIPT_DIR, "input", "Annual Summary", "2025", "2025年停电用户数据.xlsx"
)

# --- 输出根目录（日常统计 → output/daily/<日期范围>/） ---
OUTPUT_BASE_DIR      = _os.path.join(_SCRIPT_DIR, "output", "daily")
# 历史处理结果：Newdata 日期窗口内整段替换，窗口外保留
EXISTINGDATA_FOLDER  = OUTPUT_BASE_DIR

# --- 以下留空则自动命名（推荐留空） ---
RESULT_OUTPUT_PATH   = ""    # 留空 → output/daily/<日期范围>/<日期范围>停电用户_处理结果.xlsx
PUBLISH_OUTPUT_PATH  = ""    # 留空 → output/daily/<日期范围>/<日期范围>停电用户_处理结果_发出版.xlsx
STATS_OUTPUT_PATH    = ""    # 留空 → output/daily/<日期范围>/<日期范围>统计表.xlsx

# --- 表11重大事件日剔除（分年配置） ---
EXCLUDE_BASE_DIR = _os.path.join(_SCRIPT_DIR, "input", "exclude")
EXCLUDE_2025_DIR = _os.path.join(EXCLUDE_BASE_DIR, "2025")
EXCLUDE_2026_DIR = _os.path.join(EXCLUDE_BASE_DIR, "2026")

# --- 运行参数 ---
TODAY_DATE             = ""  # 今日日期，格式 YYYY-MM-DD，留空则自动取系统日期
MAJOR_EVENT_DATES_STR  = ""  # 额外 2026 重大事件停电日，逗号分隔，格式 YYYY-MM-DD
MAJOR_EVENT_FILE_PATH  = ""  # 额外 2026 重大事件日文件（txt/csv/xlsx），留空则不读取
# ============================================================

import argparse
import hashlib
import json
import os
import re
import sys
import tempfile
import time as _time_module
from collections import deque
from datetime import date, datetime, time, timedelta
from pathlib import Path
from typing import Iterable

import openpyxl
import pandas as pd

from report_io import write_tables_streaming, write_text_atomic
from postgres_source import load_active_records
from statistics_priority import (
    RULE_VERSION, FREQUENT_RULES, WARNING_RULES, classify_text,
    classify_entity_groups, official_section,
)
from report_comparison import compare_reports


RAW_ADDED_COLUMNS = ["用户停电总次数", "是否统计", "不统计原因", "频繁停电类型", "停电预警类型"]
COUNT_COLUMNS = ["停电总次数", "故障停电次数", "预安排停电次数"]
INFO_COLUMNS = [
    "用户编码",
    "中电联用户名称",
    "用户性质",
    "所属馈线编码",
    "所属馈线名称",
    "所属供电所",
    "所属区局",
    "所属地市",
]
LINE_INFO_COLUMNS = ["所属馈线编码", "所属馈线名称", "所属供电所", "所属区局", "所属地市"]
TEXT_COLUMNS = {"工单号", "用户编码", "所属馈线编码"}
NUMBER_COLUMNS = {
    "停电总次数",
    "故障停电次数",
    "预安排停电次数",
    "用户停电总次数",
}
REPEATED_LINE_COLUMNS = (
    LINE_INFO_COLUMNS
    + COUNT_COLUMNS
    + [
        "频繁停电类型",
        "停电预警类型",
        "预警时间",
        "是否频繁停",
        "是否纳入国家能源局台账",
    ]
)


def _warning_date_text(value, fallback: date | None = None) -> str:
    """Normalize 预警时间; return empty when invalid and no fallback."""
    if value is None or (isinstance(value, float) and pd.isna(value)):
        text = ""
    else:
        text = str(value).strip()
    if not text or text.lower() in {"nan", "nat", "none"}:
        return fallback.isoformat() if fallback is not None else ""
    parsed = pd.to_datetime(text, errors="coerce")
    if pd.isna(parsed):
        return fallback.isoformat() if fallback is not None else ""
    return parsed.date().isoformat()


def normalize_identifier(value) -> str:
    if pd.isna(value):
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    if isinstance(value, int):
        return str(value)
    text = str(value).strip()
    if text.endswith(".0"):
        return text[:-2]
    return text


def normalize_identifier_series(series: pd.Series) -> pd.Series:
    return series.map(normalize_identifier).astype(str)


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="按停电可靠性规则生成统计/清单 sheet。日常 Newdata 按日期窗口整段替换历史 output。"
    )
    parser.add_argument(
        "input",
        nargs="?",
        default=None,
        help="原始停电用户 Excel 文件路径（与 --newdata 二选一）",
    )
    parser.add_argument(
        "--newdata",
        default=None,
        help="新数据文件夹路径（Newdata），自动选取最新 Excel。其停电开始日闭区间整段替换历史 output，缺天不回填",
    )
    parser.add_argument(
        "--existingdata",
        default=None,
        help="历史处理结果根目录（默认 output/）。窗口外原始明细保留；窗口内以 Newdata 为准。同时作为预警记录基准",
    )
    parser.add_argument(
        "-o",
        "--output",
        help="输出 Excel 文件路径；文件夹模式下默认自动命名并保存至 output/<日期范围>/ 目录",
    )
    parser.add_argument(
        "--sheet",
        default=0,
        help="原始数据 sheet 名或序号，默认读取第 1 个 sheet（仅单文件模式生效）",
    )
    parser.add_argument(
        "--today",
        default=None,
        help="今日日期，格式 YYYY-MM-DD；不传则使用运行当天日期",
    )
    parser.add_argument(
        "--major-event-dates",
        default="",
        help="额外 2026 重大事件停电日，逗号分隔，格式 YYYY-MM-DD；与表11剔除并集合并",
    )
    parser.add_argument(
        "--major-event-file",
        default=None,
        help="额外 2026 重大事件停电日文件，支持 txt/csv/xlsx，与表11剔除并集合并",
    )
    parser.add_argument(
        "--exclude-2025-file",
        default=None,
        help="2025 年表11重大事件日剔除文件（单文件覆盖）；默认从 input/exclude/2025/ 取最新",
    )
    parser.add_argument(
        "--exclude-2025-dir",
        default=None,
        help="2025 年表11重大事件日剔除目录；默认 input/exclude/2025/，自动取最新文件",
    )
    parser.add_argument(
        "--exclude-2026-dir",
        default=None,
        help="2026 年表11重大事件日剔除目录；默认 input/exclude/2026/，自动取最新文件",
    )
    parser.add_argument(
        "--previous-output",
        default=None,
        help="上一日或历史处理结果 Excel；传入后会保留既有停电预警线路清单（累计）的预警时间，仅追加今日新增预警线路",
    )
    parser.add_argument(
        "--validate",
        action="store_true",
        help="输出规则处理摘要，便于核对剔除和清单数量",
    )
    parser.add_argument(
        "--publish-output",
        default=None,
        help="发出版 Excel 输出路径；默认在处理结果文件名后追加 _发出版.xlsx",
    )
    parser.add_argument(
        "--nea-ledger",
        default=None,
        help="国家能源局台账文件路径",
    )
    parser.add_argument(
        "--non-interactive",
        action="store_true",
        help="非交互模式：步骤失败直接退出；Newdata 并列最新时按创建时间自动选取",
    )
    parser.add_argument(
        "--period-start",
        default=None,
        help="历史区间起点 YYYY-MM-DD；与 --period-end 同时传入时启用区间过滤",
    )
    parser.add_argument(
        "--period-end",
        default=None,
        help="历史区间终点 YYYY-MM-DD；同时作为 today（截止日）",
    )
    parser.add_argument(
        "--data-year",
        type=int,
        choices=[2025, 2026],
        default=None,
        help="主数据年份策略：2025=Annual Summary；2026=Newdata 最新",
    )
    parser.add_argument(
        "--annual-2025-path",
        default=None,
        help="覆盖 2025 年年度汇总数据路径（默认 input/Annual Summary/2025/...）",
    )
    parser.add_argument(
        "--exclude-2026-file",
        default=None,
        help="2026 年表11重大事件日剔除文件（单文件覆盖）；默认从目录取最新",
    )
    return parser.parse_args()

def classify_line_category(
    row: pd.Series,
) -> str:

    # =========================
    # 最高优先级：频繁停电
    # =========================
    if str(row.get("是否频繁停", "")).strip() == "是":
        return "频繁停电"

    return classify_text(row.get("停电预警类型", ""), WARNING_RULES)


def classify_warning_category(type_str: str) -> str:
    """对用户/线路停电预警类型按优先级做互斥分类。"""
    return classify_text(type_str, WARNING_RULES)


def _stats_dict_total(count_by_city: dict[str, int]) -> int:
    return sum(int(v) for v in count_by_city.values())


def _collect_stats_section(
    count_by_city: dict[str, int],
    names_df: pd.DataFrame | None,
    title: str,
    *,
    show_line_names: bool,
) -> dict:
    """收集摘要段落结构化数据，供排版输出。"""
    total = _stats_dict_total(count_by_city)
    items: list[tuple[str, int, str]] = []
    if total > 0:
        city_items = sorted(
            count_by_city.items(),
            key=lambda item: (-int(item[1]), str(item[0])),
        )
        for city, count in city_items:
            count = int(count)
            if count <= 0:
                continue
            city_name = str(city).replace("供电局", "局").strip()
            line_names = ""
            if (
                show_line_names
                and names_df is not None
                and not names_df.empty
                and "所属地市" in names_df.columns
            ):
                city_lines = names_df[
                    names_df["所属地市"].fillna("").astype(str).eq(str(city))
                ]
                line_names = "、".join(
                    city_lines["所属馈线名称"].dropna().astype(str).unique()
                )
            items.append((city_name, count, line_names))
    return {"title": title, "total": total, "items": items}


SUMMARY_SECTION_LABELS = ("一", "二", "三", "四")


def format_summary_text(
    cutoff_text: str,
    sections: list[dict],
    *,
    empty: bool = False,
) -> str:
    """将摘要段落格式化为多行排版文本。"""
    if empty:
        return (
            f"数据截止：截至昨日（{cutoff_text}）\n\n"
            "本期无重复停电线路。\n"
        )

    lines = [f"数据截止：截至昨日（{cutoff_text}）", ""]
    for idx, section in enumerate(sections):
        label = (
            SUMMARY_SECTION_LABELS[idx]
            if idx < len(SUMMARY_SECTION_LABELS)
            else str(idx + 1)
        )
        title = section["title"]
        total = int(section["total"])
        if total == 0:
            lines.append(f"{label}、{title}：0 回")
            lines.append("")
            continue

        lines.append(f"{label}、{title}：共 {total} 回")
        for city_name, count, line_names in section["items"]:
            if line_names:
                lines.append(f"    · {city_name} {count} 回：{line_names}")
            else:
                lines.append(f"    · {city_name} {count} 回")
        lines.append("")

    lines.append("请各单位做好频繁停电管控。")
    return "\n".join(lines).rstrip() + "\n"


def summary_banner(title: str, width: int = 52) -> str:
    """生成摘要标题分隔线。"""
    line = "─" * width
    return f"{line}\n  {title}\n{line}"


def format_comparison_summary(comparison: dict) -> str:
    """Build a concise non-zero comparison sentence for report summaries."""
    if not comparison.get("available"):
        return ""

    user = comparison["metrics"]["用户"]
    line = comparison["metrics"]["线路"]
    sections: list[str] = []

    report_changes: list[str] = []
    for label, metrics, unit in (
        ("频繁停电用户", user, "户"),
        ("频繁停电线路", line, "条"),
    ):
        delta = int(metrics.get("report_total_delta", metrics.get("net_change", 0)))
        if delta < 0:
            report_changes.append(f"{label}减少 {-delta} {unit}")
        elif delta > 0:
            report_changes.append(f"{label}增加 {delta} {unit}")
    if report_changes:
        sections.append("统计表变化：" + "、".join(report_changes))

    exits: list[str] = []
    if int(user.get("left_frequent", 0)) > 0:
        exits.append(f"用户 {user['left_frequent']} 户")
    if int(line.get("left_frequent", 0)) > 0:
        exits.append(f"线路 {line['left_frequent']} 条")
    if exits:
        sections.append("实际退出频繁清单：" + "、".join(exits))

    detail_parts: list[str] = []
    for label, metrics, unit in (("用户", user, "户"), ("线路", line, "条")):
        decreased = int(metrics.get("decreased_outage_count", 0))
        type_changed = int(metrics.get("same_count_type_changed", 0))
        total = int(metrics.get("change_detail_count", decreased + type_changed))
        if total <= 0:
            continue
        breakdown: list[str] = []
        if decreased > 0:
            breakdown.append(f"停电次数下降 {decreased} {unit}")
        if type_changed > 0:
            breakdown.append(f"次数未变但类型变化 {type_changed} {unit}")
        suffix = f"（{'、'.join(breakdown)}）" if breakdown else ""
        detail_parts.append(f"{label} {total} {unit}{suffix}")
    if detail_parts:
        sections.append("变化明细：" + "，".join(detail_parts))

    if not sections:
        return "较上次成功日报无统计变化。"
    return "较上次成功日报" + "；".join(sections) + "。"


def _build_stats_section_text(
    count_by_city: dict[str, int],
    names_df: pd.DataFrame | None,
    title: str,
    *,
    show_line_names: bool,
) -> str:
    """基于统计表同口径的地市计数字典生成摘要段落（单行兼容格式）。"""
    section = _collect_stats_section(
        count_by_city, names_df, title, show_line_names=show_line_names
    )
    total = section["total"]
    if total == 0:
        return f"{title}0回；"

    bureau_parts = []
    for city_name, count, line_names in section["items"]:
        if line_names:
            bureau_parts.append(f"{city_name}{count}回：{line_names}")
        else:
            bureau_parts.append(f"{city_name}{count}回")

    bureau_text = "；".join(bureau_parts)
    return f"{title}{total}回（{bureau_text}）；"


def _prepare_warning_lines_for_summary(
    warning_lines: pd.DataFrame,
    frequent_lines_df: pd.DataFrame,
) -> pd.DataFrame:
    warning_df = warning_lines.copy() if warning_lines is not None else pd.DataFrame()
    if warning_df.empty:
        return warning_df

    warning_df["所属馈线编码"] = normalize_identifier_series(warning_df["所属馈线编码"])
    for col in ["所属地市", "所属馈线名称", "停电预警类型", "是否频繁停"]:
        if col in warning_df.columns:
            warning_df[col] = warning_df[col].fillna("").astype(str)

    frequent_codes: set[str] = set()
    if frequent_lines_df is not None and not frequent_lines_df.empty:
        frequent_codes = set(
            normalize_identifier_series(frequent_lines_df["所属馈线编码"])
            .astype(str)
            .str.strip()
        )

    warning_df = warning_df[
        ~warning_df["所属馈线编码"].astype(str).str.strip().isin(frequent_codes)
    ].copy()
    warning_df["线路分类"] = warning_df.apply(classify_line_category, axis=1)
    return warning_df


def build_summary_text(
    tables: dict[str, pd.DataFrame],
    today: date,
    stats_dicts: list[dict[str, int]],
    *,
    data_end_date: date | None = None,
    all_show_line_names: bool = False,
) -> str:
    # 摘要截止日：优先用数据最大停电日；未传入时按“运行日 yesterday”理解
    cutoff = data_end_date if data_end_date is not None else today - timedelta(days=1)
    cutoff_text = f"{cutoff.month}月{cutoff.day}日"

    (
        _wu_45, _wu_50, _wu_30, _wu_tot,
        wl_45, wl_50, wl_30, _wl_tot,
        _fu_5, _fu_60, _fu_pre, _fu_tot,
        _fl_5, _fl_60, _fl_pre, fl_tot,
        _fn_5, _fn_60, _fn_pre, _fn_tot,
    ) = stats_dicts

    frequent_df = tables.get("频繁停电线路清单", pd.DataFrame()).copy()
    if not frequent_df.empty:
        frequent_df["所属馈线编码"] = normalize_identifier_series(frequent_df["所属馈线编码"])
        frequent_df = frequent_df[
            frequent_df["频繁停电类型"].fillna("").astype(str).str.strip().ne("")
        ].copy()
        for col in ["所属地市", "所属馈线名称"]:
            frequent_df[col] = frequent_df[col].fillna("").astype(str)

    warning_df = _prepare_warning_lines_for_summary(
        tables.get("停电预警线路清单", pd.DataFrame()),
        frequent_df,
    )

    if _stats_dict_total(fl_tot) == 0 and warning_df.empty:
        return format_summary_text(cutoff_text, [], empty=True)

    show_45 = True
    show_50 = all_show_line_names
    show_30 = all_show_line_names

    sections = [
        _collect_stats_section(
            fl_tot,
            frequent_df,
            "频繁停电线路",
            show_line_names=True,
        ),
        _collect_stats_section(
            wl_45,
            warning_df[warning_df["线路分类"] == "一年4-5次"] if not warning_df.empty else warning_df,
            "年度累计停电4-5次线路",
            show_line_names=show_45,
        ),
        _collect_stats_section(
            wl_50,
            warning_df[warning_df["线路分类"] == "50天3次"] if not warning_df.empty else warning_df,
            "近50天停电3次线路",
            show_line_names=show_50,
        ),
        _collect_stats_section(
            wl_30,
            warning_df[warning_df["线路分类"] == "30天2次"] if not warning_df.empty else warning_df,
            "近30天停电2次线路",
            show_line_names=show_30,
        ),
    ]
    return format_summary_text(cutoff_text, sections)


def resolve_column(columns: Iterable[str], candidates: Iterable[str], contains: str | None = None) -> str:
    column_list = [str(c).strip() for c in columns]
    for candidate in candidates:
        if candidate in column_list:
            return candidate
    if contains:
        matches = [c for c in column_list if contains in c]
        if matches:
            return matches[0]
    raise KeyError(f"缺少必要字段：{list(candidates)}")


def normalize_sheet_arg(sheet: str) -> str | int:
    try:
        return int(sheet)
    except (TypeError, ValueError):
        return sheet


def parse_one_date(value) -> date | None:
    if pd.isna(value):
        return None
    parsed = pd.to_datetime(value, errors="coerce")
    if pd.isna(parsed):
        return None
    return parsed.date()


def parse_date_list(text: str) -> set[date]:
    result: set[date] = set()
    for item in str(text or "").replace("；", ",").replace("，", ",").split(","):
        item = item.strip()
        if not item:
            continue
        parsed = parse_one_date(item)
        if parsed:
            result.add(parsed)
    return result


def read_major_event_dates(path: str | None) -> set[date]:
    if not path:
        return set()
    p = Path(path)
    if not p.exists():
        raise FileNotFoundError(f"重大事件日期文件不存在：{p}")
    if p.suffix.lower() in {".xlsx", ".xls"}:
        frame = pd.read_excel(p, header=None)
        values = frame.to_numpy().ravel().tolist()
    else:
        values = p.read_text(encoding="utf-8-sig").replace("\n", ",").split(",")
    return {d for d in (parse_one_date(v) for v in values) if d}


def read_major_event_dates_from_table11(
    path: str | Path,
    *,
    expected_year: int | None = None,
) -> set[date]:
    """从表11重大事件日统计表读取「重大事件日期」列。"""
    p = Path(path)
    if not p.exists():
        raise FileNotFoundError(f"表11重大事件日文件不存在：{p}")
    df = pd.read_excel(p)
    col = next((c for c in df.columns if "重大事件日期" in str(c)), None)
    if col is None:
        raise KeyError(f"表11中未找到「重大事件日期」列：{p}")
    dates: set[date] = set()
    for value in df[col]:
        parsed = parse_one_date(value)
        if parsed is None:
            continue
        if expected_year is not None and parsed.year != expected_year:
            continue
        dates.add(parsed)
    return dates


_TABLE11_TIMESTAMP_RE = re.compile(r"_(\d{14,17})(?:\.xlsx)?$", re.IGNORECASE)


def _table11_filename_timestamp(path: Path) -> int | None:
    match = _TABLE11_TIMESTAMP_RE.search(path.stem)
    if not match:
        return None
    return int(match.group(1))


def select_latest_table11_file(folder: str | Path) -> Path:
    """从目录中选取最新的表11重大事件日 xlsx 文件。"""
    root = Path(folder)
    if not root.is_dir():
        raise FileNotFoundError(f"表11剔除目录不存在：{root}")

    candidates = [
        f
        for f in list(root.glob("*.xlsx")) + list(root.glob("*.xls"))
        if not f.name.startswith("~$")
    ]
    if not candidates:
        raise FileNotFoundError(f"表11剔除目录中未找到 Excel 文件：{root}")

    preferred = [
        f for f in candidates
        if "表11" in f.name or "重大事件" in f.name
    ]
    pool = preferred or candidates

    def sort_key(path: Path) -> tuple[int, float, str]:
        ts = _table11_filename_timestamp(path)
        return (
            1 if ts is not None else 0,
            float(ts) if ts is not None else path.stat().st_mtime,
            path.name,
        )

    selected = max(pool, key=sort_key)
    return selected


def load_exclude_major_dates(
    exclude_2025_dir: str,
    exclude_2026_dir: str,
    *,
    exclude_2025_file: str | None = None,
    exclude_2026_file: str | None = None,
) -> tuple[set[date], set[date], dict[str, str]]:
    """加载分年表11重大事件日剔除配置，返回 (2025日期集, 2026日期集, 来源信息)。"""
    if exclude_2025_file:
        path_2025 = Path(exclude_2025_file)
        if not path_2025.exists():
            raise FileNotFoundError(f"表11重大事件日文件不存在：{path_2025}")
    else:
        path_2025 = select_latest_table11_file(exclude_2025_dir)
    major_dates_2025 = read_major_event_dates_from_table11(path_2025, expected_year=2025)

    if exclude_2026_file:
        path_2026 = Path(exclude_2026_file)
        if not path_2026.exists():
            raise FileNotFoundError(f"表11重大事件日文件不存在：{path_2026}")
    else:
        path_2026 = select_latest_table11_file(exclude_2026_dir)
    major_dates_2026 = read_major_event_dates_from_table11(path_2026, expected_year=2026)

    source_info = {
        "2025_file": str(path_2025),
        "2026_file": str(path_2026),
    }
    return major_dates_2025, major_dates_2026, source_info


def read_previous_repeated_lines(path: str | None) -> pd.DataFrame:
    if not path:
        return pd.DataFrame(columns=REPEATED_LINE_COLUMNS)
    p = Path(path)
    if not p.exists():
        raise FileNotFoundError(f"历史处理结果文件不存在：{p}")
    try:
        previous = pd.read_excel(p, sheet_name="停电预警线路清单（累计）", dtype=str)
    except ValueError as exc:
        raise ValueError(f"历史处理结果中缺少 sheet：停电预警线路清单（累计） ({p})") from exc
    except Exception as exc:
        print(
            f"  历史预警记录文件损坏，忽略累计记录（{p.name}）：{exc}",
            file=sys.stderr,
        )
        return pd.DataFrame(columns=REPEATED_LINE_COLUMNS)
    for col in REPEATED_LINE_COLUMNS:
        if col not in previous.columns:
            previous[col] = ""
    previous = previous[REPEATED_LINE_COLUMNS].copy()
    previous["所属馈线编码"] = normalize_identifier_series(previous["所属馈线编码"])
    previous["预警时间"] = previous["预警时间"].map(
        lambda v: _warning_date_text(v) or ""
    )
    return previous

def read_nea_ledger(path: str | None) -> set[str]:

    if not path:
        return set()

    p = Path(path)

    if not p.exists():
        raise FileNotFoundError(
            f"国家能源局台账文件不存在：{p}"
        )

    df = pd.read_excel(
    p,
    dtype=str,
    keep_default_na=False,
    )

    df.columns = [
        str(c).strip()
        for c in df.columns
    ]

    gis_col = resolve_column(
        df.columns,
        ["线路GISID"],
        contains="GISID"
    )

    gis_ids = (
        df[gis_col]
        .dropna()
        .map(normalize_identifier)
        .astype(str)
    )

    return set(gis_ids)


def outage_type_mask(series: pd.Series, keyword: str) -> pd.Series:
    return series.astype(str).str.contains(keyword, na=False)


def max_count_in_rolling_window(starts: Iterable[pd.Timestamp], days: int) -> int:
    valid_starts = sorted(ts for ts in starts if pd.notna(ts))
    window: deque[pd.Timestamp] = deque()
    best = 0
    span = pd.Timedelta(days=days)
    for ts in valid_starts:
        while window and ts - window[0] >= span:
            window.popleft()
        window.append(ts)
        best = max(best, len(window))
    return best


def recent_count(
    starts: Iterable[pd.Timestamp],
    today: date,
    days: int,
) -> int:

    start_at = datetime.combine(
        today - timedelta(days=days - 1),
        time.min
    )

    end_at = datetime.combine(
        today,
        time.max
    )

    return sum(
        1
        for ts in starts
        if pd.notna(ts)
        and start_at <= ts.to_pydatetime() <= end_at
    )


def _sorted_outage_starts(starts: Iterable) -> list[pd.Timestamp]:
    parsed = pd.to_datetime(list(starts), errors="coerce")
    return sorted(ts for ts in parsed if pd.notna(ts))


def _starts_in_recent_window(
    starts: list[pd.Timestamp],
    today: date,
    days: int,
) -> list[pd.Timestamp]:
    start_at = datetime.combine(today - timedelta(days=days - 1), time.min)
    end_at = datetime.combine(today, time.max)
    return [
        ts
        for ts in starts
        if start_at <= ts.to_pydatetime() <= end_at
    ]


def compute_warning_trigger_date(
    starts: Iterable,
    today: date,
    warning_type_str: str,
) -> date | None:
    """返回触发当前预警类型的那次停电日期（非脚本运行日）。"""
    labels = [
        x.strip()
        for x in str(warning_type_str or "").split("；")
        if x.strip()
    ]
    if not labels:
        return None

    parsed = _sorted_outage_starts(starts)
    if not parsed:
        return None

    trigger_dates: list[date] = []

    if any("4-5" in label or "4次" in label for label in labels):
        if len(parsed) >= 4:
            trigger_dates.append(parsed[3].date())

    if any("50天" in label for label in labels):
        in_window = _starts_in_recent_window(parsed, today, 50)
        if len(in_window) >= 3:
            trigger_dates.append(in_window[2].date())

    if any("30天" in label for label in labels):
        in_window = _starts_in_recent_window(parsed, today, 30)
        if len(in_window) >= 2:
            trigger_dates.append(in_window[1].date())

    if trigger_dates:
        return max(trigger_dates)
    return today


def join_types(types: list[str]) -> str:

    labels = sorted(set(
        x.strip()
        for x in types
        if str(x).strip()
    ))

    return "；".join(labels)


def unique_type_labels(values: Iterable) -> list[str]:
    labels: list[str] = []
    seen: set[str] = set()
    for value in values:
        if pd.isna(value):
            continue
        for label in str(value).split("；"):
            label = label.strip()
            if label and label not in seen:
                seen.add(label)
                labels.append(label)
    return labels


def choose_info_rows(counted: pd.DataFrame, cols: dict[str, str]) -> pd.DataFrame:
    counted = counted.copy()

    info = (
        counted
        .sort_values(
            [
                cols["user"],
                cols["feeder_code"],
                cols["start"],
            ],
            kind="mergesort",
        )
        .groupby(
            [
                cols["user"],
                cols["feeder_code"],
            ],
            as_index=False,
        )
        .first()
    )
    rename_map = {
        cols["user"]: "用户编码",
        cols["name"]: "中电联用户名称",
        cols["nature"]: "用户性质",
        cols["feeder_code"]: "所属馈线编码",
        cols["feeder_name"]: "所属馈线名称",
        cols["station"]: "所属供电所",
        cols["bureau"]: "所属区局",
        cols["city"]: "所属地市",
    }
    available = [source for source in rename_map if source in info.columns]
    return info[available].rename(columns=rename_map)


def build_user_summary(
    counted: pd.DataFrame,
    cols: dict[str, str],
    today: date,
) -> pd.DataFrame:

    if counted.empty:
        return pd.DataFrame(
            columns=INFO_COLUMNS
            + COUNT_COLUMNS
            + [
                "频繁停电类型",
                "停电预警类型",
            ]
        )

    counted = counted.copy()

    # ===========================
    # 按 用户编码 + 所属馈线编码 分组
    # ===========================

    grouped = counted.groupby(
        [
            cols["user"],
            cols["feeder_code"],
        ],
        dropna=False,
    )

    total = (
        grouped.size()
        .rename("停电总次数")
    )

    fault = (
        grouped[cols["type"]]
        .apply(
            lambda s: int(
                outage_type_mask(
                    s,
                    "故障"
                ).sum()
            )
        )
        .rename("故障停电次数")
    )

    planned = (
        grouped[cols["type"]]
        .apply(
            lambda s: int(
                outage_type_mask(
                    s,
                    "预安排"
                ).sum()
            )
        )
        .rename("预安排停电次数")
    )

    starts = (
        grouped[cols["start"]]
        .apply(list)
        .rename("_starts")
    )

    summary = pd.concat(
        [
            total,
            fault,
            planned,
            starts,
        ],
        axis=1,
    ).reset_index()

    # 字段统一名称

    summary = summary.rename(
        columns={
            cols["user"]: "用户编码",
            cols["feeder_code"]: "所属馈线编码",
        }
    )

    summary["_max_60天"] = summary["_starts"].apply(
        lambda values:
        max_count_in_rolling_window(
            values,
            60,
        )
    )

    summary["_近50天"] = summary["_starts"].apply(
        lambda values:
        recent_count(
            values,
            today,
            50,
        )
    )

    summary["_近30天"] = summary["_starts"].apply(
        lambda values:
        recent_count(
            values,
            today,
            30,
        )
    )

    # ===========================
    # 频繁停电类型
    # ===========================

    def frequent_type(
        row: pd.Series,
    ) -> str:

        labels = []

        if row["停电总次数"] > 5:
            labels.append(
                "一年内停电次数超过5次"
            )

        if row["_max_60天"] > 3:
            labels.append(
                "连续60天停电次数超过3次"
            )

        if row["预安排停电次数"] > 3:
            labels.append(
                "一年内预安排停电次数超过3次"
            )

        return join_types(labels)

    # ===========================
    # 停电预警类型
    # ===========================

    def warning_type(
        row: pd.Series,
    ) -> str:

        labels = []

        if (
            4
            <= row["停电总次数"]
            <= 5
        ):
            labels.append(
                "一年内停电4-5次"
            )

        if row["_近50天"] == 3:
            labels.append(
                "近50天停电3次"
            )

        if row["_近30天"] == 2:
            labels.append(
                "近30天停电2次"
            )

        return join_types(labels)

    summary["频繁停电类型"] = summary.apply(
        frequent_type,
        axis=1,
    )

    summary["停电预警类型"] = summary.apply(
        warning_type,
        axis=1,
    )
    summary = summary.drop(columns=["_starts"])

    # ===========================
    # 合并用户基础信息
    # ===========================

    info = choose_info_rows(
        counted,
        cols,
    )

    summary = info.merge(
        summary,
        on=[
            "用户编码",
            "所属馈线编码",
        ],
        how="right",
    )

    output_cols = (
        INFO_COLUMNS
        + COUNT_COLUMNS
        + [
            "频繁停电类型",
            "停电预警类型",
        ]
    )

    return (
        summary[
            output_cols
        ]
        .sort_values(
            [
                "停电总次数",
                "用户编码",
                "所属馈线编码",
            ],
            ascending=[
                False,
                True,
                True,
            ],
        )
    )


def mark_counted_rows(raw: pd.DataFrame, cols: dict[str, str], major_dates: set[date]) -> tuple[pd.DataFrame, dict[str, int]]:
    df = raw.copy()
    df[cols["start"]] = pd.to_datetime(df[cols["start"]], errors="coerce")
    df[cols["end"]] = pd.to_datetime(df[cols["end"]], errors="coerce")

    duration_minutes = (df[cols["end"]] - df[cols["start"]]).dt.total_seconds() / 60
    major_mask = df[cols["start"]].dt.date.isin(major_dates) if major_dates else pd.Series(False, index=df.index)
    short_mask = duration_minutes < 5
    workorder_user_count = df.groupby(cols["order"])[cols["user"]].transform("nunique")
    single_user_workorder_mask = workorder_user_count <= 1

    excluded = major_mask | short_mask | single_user_workorder_mask
    counted = pd.Series(True, index=df.index)
    counted[excluded] = False
    exclude_reasons = pd.Series("", index=df.index, dtype="object")

    def add_reason(mask: pd.Series, reason: str) -> None:
        nonlocal exclude_reasons
        mask = mask.fillna(False)
        empty_mask = mask & exclude_reasons.eq("")
        filled_mask = mask & exclude_reasons.ne("")
        exclude_reasons.loc[empty_mask] = reason
        exclude_reasons.loc[filled_mask] = exclude_reasons.loc[filled_mask] + "；" + reason

    add_reason(major_mask, "重大事件停电日")
    add_reason(short_mask, "停电时长小于5分钟")
    add_reason(single_user_workorder_mask, "单一用户工单")

    six_hour_mask = pd.Series(False, index=df.index)
    df["_user_feeder_key"] = (
    normalize_identifier_series(df[cols["user"]])
    + "_"
    + normalize_identifier_series(df[cols["feeder_code"]])
    )

    for _, group in (
        df[~excluded]
        .sort_values(
            ["_user_feeder_key", cols["start"]],
            kind="mergesort"
        )
        .groupby(
            "_user_feeder_key",
            dropna=False
        )
    ):
        last_kept_start: pd.Timestamp | None = None
        for idx, row in group.iterrows():
            current_start = row[cols["start"]]
            if pd.isna(current_start):
                continue
            if last_kept_start is not None and current_start - last_kept_start < pd.Timedelta(hours=6):
                six_hour_mask.at[idx] = True
                counted.at[idx] = False
            else:
                last_kept_start = current_start

    add_reason(six_hour_mask, "同一用户6小时内多次停电合并")
    df["_是否统计_bool"] = counted
    df["_不统计原因"] = exclude_reasons.where(~counted, "")
    stats = {
        "原始行数": int(len(df)),
        "重大事件日期剔除行数": int(major_mask.sum()),
        "停电时长小于5分钟剔除行数": int(short_mask.sum()),
        "单一用户工单剔除行数": int(single_user_workorder_mask.sum()),
        "同用户6小时内合并剔除行数": int(six_hour_mask.sum()),
        "最终统计行数": int(counted.sum()),
        "最终不统计行数": int((~counted).sum()),
    }
    return df, stats


def build_raw_statistics(raw_marked: pd.DataFrame, user_summary: pd.DataFrame, cols: dict[str, str]) -> pd.DataFrame:
    df = raw_marked.drop(columns=["_是否统计_bool", "_不统计原因"]).copy()
    counts = user_summary[["用户编码", "所属馈线编码","停电总次数", "频繁停电类型", "停电预警类型"]].copy()
    counts = counts.rename(columns={"停电总次数": "用户停电总次数"})
    df["_用户馈线_key"] = (
        normalize_identifier_series(df[cols["user"]])
        + "_"
        + normalize_identifier_series(df[cols["feeder_code"]])
    )

    counts["_用户馈线_key"] = (
        normalize_identifier_series(counts["用户编码"])
        + "_"
        + normalize_identifier_series(counts["所属馈线编码"])
    )

    df = df.merge(
        counts.drop(
            columns=[
                "用户编码",
                "所属馈线编码",
            ]
        ),
        on="_用户馈线_key",
        how="left",
    ).drop(columns=["_用户馈线_key"])
    df["用户停电总次数"] = df["用户停电总次数"].fillna(0).astype(int)
    df["是否统计"] = raw_marked["_是否统计_bool"].map({True: "是", False: "否"})
    df["不统计原因"] = raw_marked["_不统计原因"].fillna("")
    for col in ["频繁停电类型", "停电预警类型"]:
        df[col] = df[col].fillna("")

    original_cols = [c for c in raw_marked.columns if c not in {"_是否统计_bool", "_不统计原因"}]
    return df[original_cols + RAW_ADDED_COLUMNS]


def build_line_list(
    user_list: pd.DataFrame,
    *,
    exclude_all_special: bool,
    include_warning_date: date | None = None,
    min_user_count: int = 1,
    frequent_user_threshold: int = 2,
) -> pd.DataFrame:
    line_columns = LINE_INFO_COLUMNS + COUNT_COLUMNS + ["频繁停电类型", "停电预警类型"]
    result_columns = (
        line_columns + ["预警时间"]
        if include_warning_date is not None
        else line_columns
    )
    if user_list.empty:
        result = pd.DataFrame(columns=result_columns)
    else:
        rows = []
        for _, group in user_list.groupby("所属馈线编码", dropna=False):
            if exclude_all_special and group["用户性质"].astype(str).eq("专用").all():
                continue
            if group["用户编码"].nunique() < min_user_count:
                continue
            sort_cols = ["停电总次数", "故障停电次数", "预安排停电次数", "用户编码"]
            representative = group.sort_values(sort_cols, ascending=[False, False, False, True]).iloc[0]
            row = {col: representative.get(col, "") for col in LINE_INFO_COLUMNS + COUNT_COLUMNS}
            frequent_users = group[
                group["频繁停电类型"].fillna("").ne("")
            ]

            if frequent_users["用户编码"].nunique() >= frequent_user_threshold:

                row["频繁停电类型"] = join_types(
                    unique_type_labels(
                        frequent_users["频繁停电类型"]
                    )
                )

            else:

                row["频繁停电类型"] = ""
            row["停电预警类型"] = join_types(unique_type_labels(group["停电预警类型"]))
            if include_warning_date is not None:
                row["预警时间"] = include_warning_date.isoformat()
            rows.append(row)
        result = pd.DataFrame(rows).reindex(columns=result_columns)

    if include_warning_date is not None:
        result["预警时间"] = (
            pd.to_datetime(result["预警时间"], errors="coerce")
            .dt.date.astype(str)
            .replace("NaT", "")
        )
        missing_warning_date = result["预警时间"].fillna("").astype(str).str.strip().eq("")
        result.loc[missing_warning_date, "预警时间"] = include_warning_date.isoformat()
        result["是否频繁停"] = result["频繁停电类型"].fillna("").ne("").map({True: "是", False: "否"})
        result = result[
            LINE_INFO_COLUMNS
            + COUNT_COLUMNS
            + ["频繁停电类型", "停电预警类型", "预警时间", "是否频繁停"]
        ]

    if result.empty:
        return result
    return result.sort_values(["停电总次数", "所属馈线编码"], ascending=[False, True])

def normalize_type_text(value: str) -> str:

    if pd.isna(value):
        return ""

    labels = [
        x.strip()
        for x in str(value).split("；")
        if x.strip()
    ]

    # 去重 + 排序
    labels = sorted(set(labels))

    return "；".join(labels)


LINE_STATE_COMPARE_COLUMNS = [
    "停电总次数",
    "故障停电次数",
    "预安排停电次数",
    "频繁停电类型",
    "停电预警类型",
    "是否频繁停",
]


def normalize_line_history(previous: pd.DataFrame | None) -> pd.DataFrame:

    if previous is None or previous.empty:
        return pd.DataFrame(columns=REPEATED_LINE_COLUMNS)

    history = previous.copy()

    for col in REPEATED_LINE_COLUMNS:
        if col not in history.columns:
            history[col] = ""

    history = history[REPEATED_LINE_COLUMNS].copy()
    history["所属馈线编码"] = normalize_identifier_series(
        history["所属馈线编码"]
    )

    return history


def latest_history_by_line(previous: pd.DataFrame | None) -> dict[str, pd.Series]:

    history = normalize_line_history(previous)

    if history.empty:
        return {}

    history = history.copy()
    history["_预警时间_dt"] = pd.to_datetime(
        history["预警时间"],
        errors="coerce"
    )

    latest = (
        history.sort_values(
            ["所属馈线编码", "_预警时间_dt"],
            ascending=[True, False],
        )
        .drop_duplicates(
            subset=["所属馈线编码"],
            keep="first",
        )
        .drop(columns=["_预警时间_dt"])
    )

    return {
        str(row["所属馈线编码"]).strip(): row
        for _, row in latest.iterrows()
    }


def line_state_changed(
    current_row: pd.Series,
    history_row: pd.Series,
) -> bool:

    for col in LINE_STATE_COMPARE_COLUMNS:

        old_value = normalize_type_text(
            history_row.get(col, "")
        )

        new_value = normalize_type_text(
            current_row.get(col, "")
        )

        if old_value != new_value:
            return True

    return False


def apply_warning_history_to_current_lines(
    current: pd.DataFrame,
    previous: pd.DataFrame | None,
    today: date,
) -> pd.DataFrame:

    result = current.copy()

    for col in REPEATED_LINE_COLUMNS:
        if col not in result.columns:
            result[col] = ""

    result = result[REPEATED_LINE_COLUMNS].copy()
    result["所属馈线编码"] = normalize_identifier_series(
        result["所属馈线编码"]
    )

    latest_history = latest_history_by_line(previous)

    for idx, current_row in result.iterrows():

        feeder_code = str(
            current_row["所属馈线编码"]
        ).strip()

        history_row = latest_history.get(feeder_code)

        current_date = _warning_date_text(current_row.get("预警时间"), today)
        if history_row is not None and not line_state_changed(
            current_row,
            history_row,
        ):
            history_date = _warning_date_text(history_row.get("预警时间"))
            result.at[idx, "预警时间"] = history_date or current_date or today.isoformat()
        else:
            result.at[idx, "预警时间"] = current_date or today.isoformat()

    return result


def count_resolved_warning_lines(
    previous: pd.DataFrame | None,
    current: pd.DataFrame,
) -> int:

    latest_history = latest_history_by_line(previous)

    if not latest_history:
        return 0

    current_codes = set(
        normalize_identifier_series(
            current["所属馈线编码"]
        ).astype(str).str.strip()
    )

    return len(
        set(latest_history.keys()) - current_codes
    )

def build_outage_line_list(counted_rows: pd.DataFrame, cols: dict[str, str], user_summary: pd.DataFrame) -> pd.DataFrame:
    line_columns = LINE_INFO_COLUMNS + COUNT_COLUMNS + ["频繁停电类型", "停电预警类型"]
    if counted_rows.empty:
        return pd.DataFrame(columns=line_columns)

    rows_with_types = counted_rows.copy()
    rows_with_types["_user_feeder_key"] = (
        normalize_identifier_series(
            rows_with_types[cols["user"]]
        )
        + "_"
        + normalize_identifier_series(
            rows_with_types[cols["feeder_code"]]
        )
    )

    types = user_summary[
        [
            "用户编码",
            "所属馈线编码",
            "频繁停电类型",
            "停电预警类型",
        ]
    ].copy()

    types["_user_feeder_key"] = (
        normalize_identifier_series(
            types["用户编码"]
        )
        + "_"
        + normalize_identifier_series(
            types["所属馈线编码"]
        )
    )

    rows_with_types = rows_with_types.merge(
        types.drop(
            columns=[
                "用户编码",
                "所属馈线编码",
            ]
        ),
        on="_user_feeder_key",
        how="left",
    )

    rows = []
    for _, group in rows_with_types.groupby(cols["feeder_code"], dropna=False):
        representative = group.sort_values(cols["start"], kind="mergesort").iloc[0]
        row = {
            "所属馈线编码": representative.get(cols["feeder_code"], ""),
            "所属馈线名称": representative.get(cols["feeder_name"], ""),
            "所属供电所": representative.get(cols["station"], ""),
            "所属区局": representative.get(cols["bureau"], ""),
            "所属地市": representative.get(cols["city"], ""),
            "停电总次数": int(len(group)),
            "故障停电次数": int(outage_type_mask(group[cols["type"]], "故障").sum()),
            "预安排停电次数": int(outage_type_mask(group[cols["type"]], "预安排").sum()),
            "频繁停电类型": join_types(unique_type_labels(group["频繁停电类型"])),
            "停电预警类型": join_types(unique_type_labels(group["停电预警类型"])),
        }
        rows.append(row)

    result = pd.DataFrame(rows, columns=line_columns)
    return result.sort_values(["停电总次数", "所属馈线编码"], ascending=[False, True])

def add_nea_flag(
    df: pd.DataFrame,
    nea_gis_ids: set[str],
) -> pd.DataFrame:

    if df.empty:
        df["是否纳入国家能源局台账"] = ""
        return df

    result = df.copy()

    feeder_codes = (
        result["所属馈线编码"]
        .map(normalize_identifier)
        .astype(str)
    )

    result["是否纳入国家能源局台账"] = feeder_codes.isin(
        nea_gis_ids
    ).map({
        True: "是",
        False: "否"
    })

    return result

def merge_repeated_lines(
    current: pd.DataFrame,
    previous: pd.DataFrame | None,
    today: date,
) -> tuple[pd.DataFrame, int]:

    current = current.copy()

    for col in REPEATED_LINE_COLUMNS:
        if col not in current.columns:
            current[col] = ""

    current = current[REPEATED_LINE_COLUMNS]

    current["所属馈线编码"] = normalize_identifier_series(
        current["所属馈线编码"]
    )

    # =====================================
    # 无历史
    # =====================================
    if previous is None or previous.empty:

        return current, len(current)

    previous = normalize_line_history(previous)

    # =====================================
    # 最终结果 = 历史全保留
    # =====================================
    final_rows = previous.to_dict("records") if previous is not None else []

    updated_count = 0
    latest_history = latest_history_by_line(previous)

    # =====================================
    # 当前线路逐条比较
    # =====================================
    for _, current_row in current.iterrows():

        feeder_code = str(
            current_row["所属馈线编码"]
        ).strip()

        latest_row = latest_history.get(feeder_code)

        # =====================================
        # 新线路
        # =====================================
        if latest_row is None:

            final_rows.append(
                current_row.to_dict()
            )

            updated_count += 1

            continue

        # =====================================
        # 比较是否变化
        # =====================================
        changed = line_state_changed(
            current_row,
            latest_row,
        )

        # =====================================
        # 有变化 -> 新增一条
        # =====================================
        if changed:

            new_row = current_row.to_dict()
            new_row["预警时间"] = (
                _warning_date_text(new_row.get("预警时间"), today) or today.isoformat()
            )

            final_rows.append(new_row)

            updated_count += 1

    # =====================================
    # 合并
    # =====================================
    merged = pd.DataFrame(
        final_rows,
        columns=REPEATED_LINE_COLUMNS
    )

    merged = merged.sort_values(
        ["所属馈线编码", "预警时间"],
        ascending=[True, False]
    )

    return merged, updated_count

def build_tables(
    raw: pd.DataFrame,
    today: date,
    major_dates: set[date],
    previous_repeated_lines: pd.DataFrame | None = None,
    nea_gis_ids: set[str] | None = None,
) -> tuple[dict[str, pd.DataFrame], dict[str, int]]:
    raw.columns = [str(c).strip() for c in raw.columns]
    cols = {
        "order": resolve_column(raw.columns, ["工单号"]),
        "user": resolve_column(raw.columns, ["用户编码"]),
        "name": resolve_column(raw.columns, ["中电联用户名称", "用户名称"], contains="用户名称"),
        "nature": resolve_column(raw.columns, ["用户性质"]),
        "type": resolve_column(raw.columns, ["停电类型"]),
        "start": resolve_column(raw.columns, ["去重后停电开始时间", "停电开始时间"], contains="停电开始时间"),
        "end": resolve_column(raw.columns, ["去重后停电结束时间", "停电结束时间"], contains="停电结束时间"),
        "feeder_code": resolve_column(raw.columns, ["所属馈线编码"], contains="馈线编码"),
        "feeder_name": resolve_column(raw.columns, ["所属馈线名称"], contains="馈线名称"),
        "station": resolve_column(raw.columns, ["所属供电所"], contains="供电所"),
        "bureau": resolve_column(raw.columns, ["所属区局"], contains="区局"),
        "city": resolve_column(raw.columns, ["所属地市"], contains="地市"),
    }
    for key in ["order", "user", "feeder_code"]:
        raw[cols[key]] = normalize_identifier_series(raw[cols[key]])

    marked, stats = mark_counted_rows(raw, cols, major_dates)
    counted_rows = marked[marked["_是否统计_bool"]].copy()
    user_summary = build_user_summary(counted_rows, cols, today)
    user_export_columns = (
        INFO_COLUMNS
        + COUNT_COLUMNS
        + ["频繁停电类型", "停电预警类型"]
    )

    raw_statistics = build_raw_statistics(marked, user_summary, cols)
    frequent_users = user_summary[user_summary["频繁停电类型"].ne("")].copy()
    warning_users = user_summary[user_summary["停电预警类型"].ne("")].copy()
    outage_users = user_summary[user_summary["停电总次数"].ge(1)].copy()
    outage_lines = build_line_list(outage_users, exclude_all_special=True)

    # 频繁停电线路清单：从全部停电用户取代表，再按频繁停电条件过滤
    frequent_lines = build_line_list(
        outage_users,
        exclude_all_special=True,
        min_user_count=2,
        frequent_user_threshold=2,
    )
    if not frequent_lines.empty:
        frequent_lines = frequent_lines[
            frequent_lines["频繁停电类型"].fillna("").astype(str).str.strip().ne("")
        ].copy()
    frequent_line_codes = set(
        frequent_lines["所属馈线编码"]
        .astype(str)
        .str.strip()
    )

    # 停电预警线路清单：含全专用馈线，专用用户预警也需纳入线路清单
    current_repeated_lines = build_line_list(
        outage_users,
        exclude_all_special=False,
        include_warning_date=today,
    )
    if not current_repeated_lines.empty:
        current_repeated_lines = current_repeated_lines[
            current_repeated_lines["停电预警类型"].fillna("").ne("")
        ].copy()

    # =====================================
    # 只有频繁停电线路才保留频繁停电标签
    # =====================================

    mask = (
        current_repeated_lines["所属馈线编码"]
        .astype(str)
        .str.strip()
        .isin(frequent_line_codes)
    )

    # 不属于频繁停电线路
    current_repeated_lines.loc[
        ~mask,
        "频繁停电类型"
    ] = ""

    current_repeated_lines.loc[
        ~mask,
        "是否频繁停"
    ] = "否"

    # 属于频繁停电线路
    current_repeated_lines.loc[
        mask,
        "是否频繁停"
    ] = "是"
    # 当前清单：状态未变沿用累计表预警时间；新增或状态变化时用数据截止日
    warning_lines = apply_warning_history_to_current_lines(
        current_repeated_lines,
        previous_repeated_lines,
        today,
    )
    resolved_warning_line_count = count_resolved_warning_lines(
        previous_repeated_lines,
        current_repeated_lines,
    )
    repeated_lines, new_repeated_line_count = merge_repeated_lines(
        warning_lines,
        previous_repeated_lines,
        today,
    )

    # =====================================
    # 增加国家能源局台账标识
    # =====================================

    nea_gis_ids = nea_gis_ids or set()

    line_tables = {
        "停电线路清单": outage_lines,
        "停电预警线路清单": warning_lines,
        "频繁停电线路清单": frequent_lines,
        "停电预警线路清单（累计）": repeated_lines,
    }

    for name, table in line_tables.items():

        line_tables[name] = add_nea_flag(
            table,
            nea_gis_ids
        )

    outage_lines = line_tables["停电线路清单"]
    warning_lines = line_tables["停电预警线路清单"]
    frequent_lines = line_tables["频繁停电线路清单"]
    repeated_lines = line_tables["停电预警线路清单（累计）"]

    stats.update(
        {
            "统计用户数": int(user_summary[["用户编码", "所属馈线编码"]].drop_duplicates().shape[0]) if not user_summary.empty else 0,
            "频繁停电用户数": int(len(frequent_users)),
            "频繁停电线路数": int(len(frequent_lines)),
            "停电预警用户数": int(len(warning_users)),
            "停电预警线路数": int(len(warning_lines)),
            "停电线路数": int(len(outage_lines)),
            "本次全量重复停电线路数": int(len(current_repeated_lines)),
            "历史重复停电线路数": int(len(previous_repeated_lines)) if previous_repeated_lines is not None else 0,
            "历史已解除预警线路数": int(resolved_warning_line_count),
            "今日新增重复停电线路数": int(new_repeated_line_count),
            "停电预警线路清单（累计）": int(len(repeated_lines)),
        }
    )
    user_tables = {
        "停电用户清单": outage_users,
        "停电预警用户清单": warning_users,
        "频繁停电用户清单": frequent_users,
    }
    for name, table in user_tables.items():
        user_tables[name] = table.reindex(columns=user_export_columns)

    return {
        "用户停电总次数统计表": raw_statistics,
        "停电用户清单": user_tables["停电用户清单"],
        "停电线路清单": outage_lines,
        "停电预警用户清单": user_tables["停电预警用户清单"],
        "停电预警线路清单": warning_lines,
        "频繁停电用户清单": user_tables["频繁停电用户清单"],
        "频繁停电线路清单": frequent_lines,
        "停电预警线路清单（累计）": repeated_lines,
    }, stats


def format_worksheet(ws) -> None:

    ws.freeze_panes = "A2"

    ws.auto_filter.ref = ws.dimensions

    # =====================================
    # 获取文本列
    # =====================================
    text_column_indexes = [
        cell.column
        for cell in ws[1]
        if cell.value in TEXT_COLUMNS
        or str(cell.value).endswith("编码")
    ]

    # =====================================
    # 获取数值列
    # =====================================
    number_column_indexes = [
        cell.column
        for cell in ws[1]
        if cell.value in NUMBER_COLUMNS
    ]

    # =====================================
    # 文本列处理
    # =====================================
    for column_idx in text_column_indexes:

        for row in range(2, ws.max_row + 1):

            cell = ws.cell(row=row, column=column_idx)

            if cell.value is not None:

                cell.value = normalize_identifier(cell.value)

            # 强制文本
            cell.number_format = "@"

    # =====================================
    # 数值列处理
    # =====================================
    for column_idx in number_column_indexes:

        for row in range(2, ws.max_row + 1):

            cell = ws.cell(row=row, column=column_idx)

            value = cell.value

            if value is None or value == "":
                continue

            try:

                # 转整数
                cell.value = int(float(value))

            except Exception:
                pass

            # 设置数值格式
            cell.number_format = "0"

    # =====================================
    # 自动列宽
    # =====================================
    for column_cells in ws.columns:

        max_length = 0

        column_letter = column_cells[0].column_letter

        for cell in column_cells[:2000]:

            if cell.value is None:
                continue

            max_length = max(
                max_length,
                len(str(cell.value))
            )

        ws.column_dimensions[column_letter].width = min(
            max(max_length + 2, 10),
            38
        )


def write_tables(tables: dict[str, pd.DataFrame], output_path: Path) -> None:
    write_tables_streaming(
        tables,
        output_path,
        text_columns=TEXT_COLUMNS,
        number_columns=NUMBER_COLUMNS,
    )


def build_publish_tables(tables: dict[str, pd.DataFrame]) -> dict[str, pd.DataFrame]:
    raw_statistics = tables["用户停电总次数统计表"].copy()
    if "是否统计" in raw_statistics.columns:
        raw_statistics = raw_statistics[
            raw_statistics["是否统计"].astype(str).str.strip().eq("是")
        ].copy()

    published = {
        "用户停电总次数统计表": raw_statistics,
        "停电预警线路清单": tables["停电预警线路清单"].copy(),
        "停电预警线路清单（累计）": tables["停电预警线路清单（累计）"].copy(),
    }
    if "变化明细" in tables:
        published["变化明细"] = tables["变化明细"].copy()
    return published


def default_publish_output_path(output_path: Path) -> Path:
    return output_path.with_name(
        f"{output_path.stem}_发出版{output_path.suffix}"
    )


# ============================================================
# 步骤执行器：实时进度反馈 + 失败重试
# ============================================================

class StepRunner:
    """封装每个步骤的执行、计时、进度打印和失败重试逻辑。"""

    def __init__(self, *, non_interactive: bool = False) -> None:
        self.step_num = 0
        self.failed: list[str] = []
        self.non_interactive = non_interactive

    def run(self, name: str, func, *args, allow_skip: bool = False, **kwargs):
        """
        执行一个步骤。
        - 打印步骤标题和耗时。
        - 失败时提示用户选择重试 / 跳过 / 退出。
        - allow_skip=True 时允许跳过（返回 None）；否则必须重试或退出。
        """
        self.step_num += 1
        label = f"步骤 {self.step_num}：{name}"

        while True:
            _sep = "=" * 62
            print(f"\n{_sep}")
            print(f">>  {label}")
            print(_sep)
            t0 = _time_module.time()

            try:
                result = func(*args, **kwargs)
                elapsed = _time_module.time() - t0
                print(f"[OK]  {label} 完成（耗时 {elapsed:.1f}s）")
                return result

            except Exception as exc:
                elapsed = _time_module.time() - t0
                self.failed.append(name)
                print(f"\n[ERR]  {label} 失败（耗时 {elapsed:.1f}s）")
                print(f"   错误类型：{type(exc).__name__}")
                print(f"   错误详情：{exc}")

                if self.non_interactive:
                    raise exc

                while True:
                    if allow_skip:
                        prompt = "   选择操作 [y=重新运行 / s=跳过此步骤 / q=退出]: "
                    else:
                        prompt = "   选择操作 [y=重新运行 / q=退出]: "
                    choice = input(prompt).strip().lower()

                    if choice == "y":
                        print("   -> 正在重新运行...\n")
                        break  # 重试
                    elif choice == "s" and allow_skip:
                        print("   -> 跳过此步骤，继续后续处理")
                        return None
                    elif choice == "q":
                        print("   -> 用户选择退出")
                        raise SystemExit(1)
                    else:
                        print("   请输入有效选项")

    def summary(self) -> None:
        if self.failed:
            print(f"\n⚠  以下 {len(self.failed)} 个步骤失败或被跳过：")
            for name in self.failed:
                print(f"   - {name}")
        else:
            print("\n[OK]  全部步骤执行完毕，无失败步骤。")


# ============================================================
# 2025 年统计数据
# ============================================================

def read_2025_raw(file_path: str) -> pd.DataFrame:
    """从固定路径读取 2025 年原始停电数据（取列数最多的 Sheet）。"""
    target = Path(file_path)
    if not target.exists():
        raise FileNotFoundError(f"2025 年数据文件不存在：{target}")

    print(f"  读取 2025 年数据文件：{target.name}")
    with pd.ExcelFile(target) as xf:
        best_sheet, best_cols = xf.sheet_names[0], 0
        for sheet_name in xf.sheet_names:
            try:
                preview = pd.read_excel(xf, sheet_name=sheet_name, nrows=2, dtype=str)
                if len(preview.columns) > best_cols:
                    best_cols = len(preview.columns)
                    best_sheet = sheet_name
            except Exception:
                pass
        print(f"  读取 Sheet：{best_sheet}（{best_cols} 列）")
        df = pd.read_excel(xf, sheet_name=best_sheet, dtype=str)
    return df.dropna(how="all").reset_index(drop=True)


def read_2025_raw_cached(file_path: str) -> pd.DataFrame:
    """Parse the 2025 baseline once, then reuse an atomically written cache."""
    target = Path(file_path)
    if not target.exists():
        raise FileNotFoundError(f"2025 年数据文件不存在：{target}")
    digest = hashlib.sha256()
    with target.open("rb") as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(block)
    cache_root = Path(
        os.environ.get(
            "OUTAGE_CACHE_DIR",
            str(Path(OUTPUT_BASE_DIR).parent / ".cache"),
        )
    )
    cache_root.mkdir(parents=True, exist_ok=True)
    cache_path = cache_root / f"2025-normalized-v108-{digest.hexdigest()[:20]}.pkl"
    if cache_path.is_file():
        started = _time_module.perf_counter()
        cached = pd.read_pickle(cache_path)
        print(
            f"  读取 2025 基准缓存：{cache_path.name}（{len(cached)} 行，"
            f"{_time_module.perf_counter() - started:.1f}s）"
        )
        return cached

    raw = normalize_2025_columns(read_2025_raw(str(target)))
    handle = tempfile.NamedTemporaryFile(
        prefix=f".{cache_path.stem}.", suffix=".tmp", dir=cache_root, delete=False
    )
    temp_path = Path(handle.name)
    handle.close()
    try:
        raw.to_pickle(temp_path)
        os.replace(temp_path, cache_path)
    except BaseException:
        temp_path.unlink(missing_ok=True)
        raise
    # Keep only a few old source revisions to bound disk usage.
    old_caches = sorted(
        cache_root.glob("2025-normalized-v108-*.pkl"),
        key=lambda path: path.stat().st_mtime,
        reverse=True,
    )
    for old_path in old_caches[4:]:
        old_path.unlink(missing_ok=True)
    print(f"  已生成 2025 基准缓存：{cache_path.name}（{len(raw)} 行）")
    return raw


def normalize_2025_columns(df: pd.DataFrame) -> pd.DataFrame:
    """
    将 2025 年原始停电数据列名映射为本脚本标准列名。
    申报用户名称→中电联用户名称 / 用户类型→用户性质 /
    线路GISID→所属馈线编码 / 线路段名称→所属馈线名称
    同时仅保留「是否为去重用户=是」的有效记录。
    """
    df = df.copy()
    df.columns = [str(c).strip() for c in df.columns]

    rename_map = {
        "申报用户名称": "中电联用户名称",
        "用户类型": "用户性质",
        "线路GISID": "所属馈线编码",
        "线路段名称": "所属馈线名称",
    }
    actual = {
        source: target
        for source, target in rename_map.items()
        if source in df.columns and target not in df.columns
    }
    if actual:
        df = df.rename(columns=actual)

    dedup_col = next(
        (col for col in df.columns if "是否为去重用户" in col),
        None,
    )
    if dedup_col:
        before = len(df)
        df = df[df[dedup_col].astype(str).str.strip().eq("是")]
        print(f"  保留去重后用户：{before} → {len(df)} 行")

    return df.reset_index(drop=True)


def filter_raw_by_outage_date(
    raw: pd.DataFrame,
    start: date,
    end: date,
) -> pd.DataFrame:
    """按停电开始时间截取 [start, end] 区间内的原始记录。"""
    raw = raw.copy()
    raw.columns = [str(c).strip() for c in raw.columns]

    start_col = resolve_column(
        raw.columns,
        ["去重后停电开始时间", "停电开始时间"],
        contains="停电开始时间",
    )
    ts = pd.to_datetime(raw[start_col], errors="coerce")
    mask = (ts.dt.date >= start) & (ts.dt.date <= end)
    return raw[mask].reset_index(drop=True)


def prior_year_same_day(value: date) -> date:
    """将日期映射到去年同期（2 月 29 日回退到 2 月 28 日）。"""
    shifted = _safe_date(value.year - 1, value.month, value.day)
    if shifted is None:
        raise ValueError(f"无法计算去年同期日期：{value.isoformat()}")
    return shifted


def shift_dates_to_prior_year(dates: set[date]) -> set[date]:
    """将重大事件日期整体平移到上一年。"""
    shifted: set[date] = set()
    for item in dates:
        prior = _safe_date(item.year - 1, item.month, item.day)
        if prior is not None:
            shifted.add(prior)
    return shifted


def build_2025_statistics_tables(
    raw_2025_path: str,
    period_start: date,
    period_end: date,
    major_dates_2025: set[date],
    nea_gis_ids: set[str],
) -> dict[str, pd.DataFrame]:
    """
    基于 2025 年原始数据，按本次输出区间去年同期跑可靠性规则，
    生成 2025 年统计表所需的频繁停电清单。
    """
    if os.environ.get("DATA_BACKEND", "excel").strip().lower() == "postgres":
        raw = load_active_records(date(2025, 1, 1), date(2025, 12, 31))
        print(f"  从 PostgreSQL 读取 2025 年有效数据：{len(raw)} 行")
        raw = normalize_2025_columns(raw)
    else:
        raw = read_2025_raw_cached(raw_2025_path)

    yoy_start = prior_year_same_day(period_start)
    yoy_end = prior_year_same_day(period_end)
    raw = filter_raw_by_outage_date(raw, yoy_start, yoy_end)
    print(
        f"  2025年统计取值区间：{yoy_start.isoformat()} ~ {yoy_end.isoformat()}（{len(raw)} 行）"
    )

    if raw.empty:
        empty_tables, _ = build_tables(
            raw,
            yoy_end,
            major_dates_2025,
            previous_repeated_lines=pd.DataFrame(columns=REPEATED_LINE_COLUMNS),
            nea_gis_ids=nea_gis_ids,
        )
        return empty_tables

    tables, _ = build_tables(
        raw,
        yoy_end,
        major_dates_2025,
        previous_repeated_lines=pd.DataFrame(columns=REPEATED_LINE_COLUMNS),
        nea_gis_ids=nea_gis_ids,
    )
    return tables


# ============================================================
# 统计表生成
# ============================================================

def _normalize_city(name: str) -> str:
    """标准化城市名称，去除"供电局"等后缀，便于匹配。"""
    s = str(name).strip()
    for suffix in ["供电局", "供电公司", "供电所"]:
        s = s.replace(suffix, "")
    return s.strip()


def _count_by_city(
    df: pd.DataFrame,
    city_col: str,
    type_col: str | None,
    keyword: str | None = None,
    nea_only: bool = False,
) -> dict[str, int]:
    """
    按城市统计行数。
    keyword: 若不为 None，仅统计 type_col 含该关键词的行。
    nea_only: 若为 True，仅统计"是否纳入国家能源局台账=是"的行。
    """
    if df.empty or city_col not in df.columns:
        return {}
    sub = df.copy()
    if nea_only and "是否纳入国家能源局台账" in sub.columns:
        sub = sub[sub["是否纳入国家能源局台账"].astype(str).str.strip().eq("是")]
    if keyword and type_col and type_col in sub.columns:
        sub = sub[sub[type_col].fillna("").astype(str).str.contains(keyword, na=False, regex=False)]
    if sub.empty:
        return {}
    return sub.groupby(city_col).size().to_dict()


def _count_by_city_exclusive(
    df: pd.DataFrame,
    city_col: str,
    category_fn,
    category_label: str,
    *,
    nea_only: bool = False,
) -> dict[str, int]:
    """按城市统计互斥分类后的行数，每条记录只计入一个类别。"""
    if df.empty or city_col not in df.columns:
        return {}

    sub = df.copy()
    if nea_only and "是否纳入国家能源局台账" in sub.columns:
        sub = sub[sub["是否纳入国家能源局台账"].astype(str).str.strip().eq("是")]
    if sub.empty:
        return {}

    sub["_category"] = sub.apply(category_fn, axis=1)
    sub = sub[sub["_category"] == category_label]
    if sub.empty:
        return {}
    return sub.groupby(city_col).size().to_dict()


def _sum_city_dicts(*dicts: dict[str, int]) -> dict[str, int]:
    """按城市合并多个计数字典，用于合计列。"""
    result: dict[str, int] = {}
    for d in dicts:
        for city, count in d.items():
            result[city] = result.get(city, 0) + int(count)
    return result


def _match_city(template_name: str, data_cities: list[str]) -> str | None:
    """将模板短城市名与数据完整城市名进行匹配，依次尝试精确/包含/前缀匹配。"""
    t = _normalize_city(template_name)
    if not t:
        return None
    norm_map = {dc: _normalize_city(dc) for dc in data_cities}
    for dc, nd in norm_map.items():
        if nd == t:
            return dc
    for dc in data_cities:
        if t in dc:
            return dc
    if len(t) >= 2:
        for dc in data_cities:
            if dc.startswith(t[:2]):
                return dc
    return None


def _get_template_city_rows(ws) -> list[tuple[int, str]]:
    """读取模板 A 列（第3行起）的数据行，返回 [(行号, 城市名), ...]。"""
    rows = []
    for r in range(3, ws.max_row + 1):
        v = ws.cell(r, 1).value
        if v is not None and str(v).strip():
            rows.append((r, str(v).strip()))
    return rows


def _fill_stats_sheet(
    ws,
    dicts: list[dict[str, int]],
    all_cities: list[str],
    company_keyword: str = "公司",
) -> None:
    """
    向统计表 Sheet 填充数据。
    dicts: 每个列（B列起）对应的 {city: count} 字典。
    公司行取所有非公司行的列合计。
    """
    city_rows = _get_template_city_rows(ws)
    company_row: int | None = None

    for r_idx, city_name in city_rows:
        if company_keyword in city_name:
            company_row = r_idx
            continue
        matched = _match_city(city_name, all_cities)
        for col_offset, d in enumerate(dicts):
            val = int(d.get(matched, 0)) if matched else 0
            ws.cell(r_idx, col_offset + 2).value = val if val else None

    if company_row is not None:
        non_company = [r for r, cn in city_rows if company_keyword not in cn]
        for col_offset in range(len(dicts)):
            col_idx = col_offset + 2
            total = sum(ws.cell(r, col_idx).value or 0 for r in non_company)
            ws.cell(company_row, col_idx).value = total if total else None


def format_stats_workbook(wb) -> None:
    """为统计表所有单元格设置微软雅黑字体和居中对齐。"""
    from openpyxl.styles import Alignment, Font

    center_font = Font(name="微软雅黑")
    center_align = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                cell.font = center_font
                cell.alignment = center_align


def save_openpyxl_workbook(wb, output_path: Path) -> Path:
    """
    保存 openpyxl 工作簿。
    若目标文件正被 Excel 占用，则写入同目录下带时间戳的备用文件，避免步骤失败。
    """
    output_path = Path(output_path)
    os.makedirs(output_path.parent, exist_ok=True)

    try:
        wb.save(str(output_path))
        return output_path
    except PermissionError:
        pass

    fd, tmp_name = tempfile.mkstemp(
        suffix=".xlsx",
        prefix=f".{output_path.stem}_",
        dir=output_path.parent,
    )
    os.close(fd)
    tmp_path = Path(tmp_name)
    try:
        wb.save(str(tmp_path))
        try:
            os.replace(tmp_path, output_path)
            return output_path
        except PermissionError:
            fallback = output_path.with_name(
                f"{output_path.stem}_{datetime.now():%Y%m%d_%H%M%S}{output_path.suffix}"
            )
            os.replace(tmp_path, fallback)
            print(
                f"  ⚠ 无法覆盖已打开的文件：{output_path}",
                file=sys.stderr,
            )
            print(f"  统计表已改存至：{fallback}")
            print(
                "  提示：请关闭 Excel 中打开的统计表后，可将备用文件重命名替换。",
                file=sys.stderr,
            )
            return fallback
    except Exception:
        if tmp_path.exists():
            tmp_path.unlink(missing_ok=True)
        raise


def _warning_line_category_row(row: pd.Series) -> str:
    """线路预警互斥分类，排除频繁停电线路。"""
    if str(row.get("是否频繁停", "")).strip() == "是":
        return ""
    return classify_line_category(row)


def _warning_user_category_row(row: pd.Series) -> str:
    """用户预警互斥分类。"""
    return classify_warning_category(row.get("停电预警类型", ""))


def compute_2026_stats_dicts(tables_2026: dict[str, pd.DataFrame]) -> list[dict[str, int]]:
    """计算 2026 年统计表 B~U 列对应的地市计数字典，合计列均为子列之和。"""
    user_groups = classify_entity_groups(
        tables_2026["频繁停电用户清单"], tables_2026["停电预警用户清单"],
        key_column="用户编码",
    )
    line_groups = classify_entity_groups(
        tables_2026["频繁停电线路清单"], tables_2026["停电预警线路清单"],
        key_column="所属馈线编码",
    )
    nea_groups = classify_entity_groups(
        tables_2026["频繁停电线路清单"], tables_2026["停电预警线路清单"],
        key_column="所属馈线编码", nea_only=True,
    )
    warning_labels = tuple(name for _, name in reversed(WARNING_RULES))
    frequent_labels = tuple(name for _, name in FREQUENT_RULES)
    return (
        official_section(user_groups, "warning", warning_labels)
        + official_section(line_groups, "warning", warning_labels)
        + official_section(user_groups, "frequent", frequent_labels)
        + official_section(line_groups, "frequent", frequent_labels)
        + official_section(nea_groups, "frequent", frequent_labels)
    )


def compute_2025_stats_dicts(tables_2025: dict[str, pd.DataFrame]) -> list[dict[str, int]]:
    """计算 2025 年统计表 B~M 列对应的地市计数字典，合计列均为子列之和。"""
    empty = pd.DataFrame()
    labels = tuple(name for _, name in FREQUENT_RULES)
    user_groups = classify_entity_groups(tables_2025["频繁停电用户清单"], empty, key_column="用户编码")
    line_groups = classify_entity_groups(tables_2025["频繁停电线路清单"], empty, key_column="所属馈线编码")
    nea_groups = classify_entity_groups(tables_2025["频繁停电线路清单"], empty, key_column="所属馈线编码", nea_only=True)
    return (
        official_section(user_groups, "frequent", labels)
        + official_section(line_groups, "frequent", labels)
        + official_section(nea_groups, "frequent", labels)
    )


def generate_statistics_table(
    tables_2026: dict[str, pd.DataFrame],
    template_path: str,
    output_path: Path,
    *,
    raw_2025_path: str,
    period_start: date | None,
    period_end: date | None,
    major_dates_2026: set[date],
    major_dates_2025: set[date],
    nea_gis_ids: set[str],
) -> None:
    """
    生成汇总统计表：填充 2026 年数据；2025 年按去年同期区间从原始数据计算。
    """
    wb = openpyxl.load_workbook(template_path)

    dicts_2026 = compute_2026_stats_dicts(tables_2026)
    all_2026_cities = list({k for d in dicts_2026 for k in d})

    if "2026年统计表" in wb.sheetnames:
        _fill_stats_sheet(wb["2026年统计表"], dicts_2026, all_2026_cities)
        print("  2026年统计表填充完成")
    else:
        print("  ⚠ 模板中未找到 '2026年统计表' Sheet，跳过", file=sys.stderr)

    if "2025年统计表" in wb.sheetnames:
        if period_end is None:
            raise ValueError("无法确定本次数据截止日期，无法计算 2025 年统计表")
        stats_period_start = period_start or date(period_end.year, 1, 1)
        tables_2025 = build_2025_statistics_tables(
            raw_2025_path=raw_2025_path,
            period_start=stats_period_start,
            period_end=period_end,
            major_dates_2025=major_dates_2025,
            nea_gis_ids=nea_gis_ids,
        )
        dicts_2025 = compute_2025_stats_dicts(tables_2025)
        all_2025_cities = list({k for d in dicts_2025 for k in d})
        _fill_stats_sheet(wb["2025年统计表"], dicts_2025, all_2025_cities)
        yoy_end = prior_year_same_day(period_end)
        print(f"  2025年统计表填充完成（取值截止：{yoy_end.isoformat()}）")
    else:
        print("  ⚠ 模板中未找到 '2025年统计表' Sheet", file=sys.stderr)

    format_stats_workbook(wb)
    saved_path = save_openpyxl_workbook(wb, output_path)
    print(f"  统计表已保存：{saved_path}")


# ============================================================
# 文件夹模式辅助函数
# ============================================================

def _safe_date(year: int, month: int, day: int) -> date | None:
    try:
        return date(year, month, day)
    except ValueError:
        return None


def parse_dates_from_filename(name: str) -> list[date]:
    """从文件名解析所有可能的日期，供取最大值作为文件名日期。"""
    stem = Path(name).stem
    dates: list[date] = []

    for match in re.finditer(r"(\d{4})-(\d{1,2})-(\d{1,2})", stem):
        parsed = _safe_date(
            int(match.group(1)),
            int(match.group(2)),
            int(match.group(3)),
        )
        if parsed:
            dates.append(parsed)

    range_full = re.search(
        r"(\d{4})\.(\d{1,2})\.(\d{1,2})-(\d{1,2})\.(\d{1,2})",
        stem,
    )
    if range_full:
        year = int(range_full.group(1))
        start = _safe_date(year, int(range_full.group(2)), int(range_full.group(3)))
        end = _safe_date(year, int(range_full.group(4)), int(range_full.group(5)))
        if start:
            dates.append(start)
        if end:
            dates.append(end)
    else:
        range_same_month = re.search(
            r"(\d{4})\.(\d{1,2})\.(\d{1,2})-(\d{1,2})(?!\.(\d))", stem
        )
        if range_same_month:
            year = int(range_same_month.group(1))
            month = int(range_same_month.group(2))
            start = _safe_date(year, month, int(range_same_month.group(3)))
            end = _safe_date(year, month, int(range_same_month.group(4)))
            if start:
                dates.append(start)
            if end:
                dates.append(end)
        else:
            for match in re.finditer(r"(\d{4})\.(\d{1,2})\.(\d{1,2})", stem):
                parsed = _safe_date(
                    int(match.group(1)),
                    int(match.group(2)),
                    int(match.group(3)),
                )
                if parsed:
                    dates.append(parsed)

    return dates


def get_newdata_effective_date(path: Path) -> tuple[date, datetime, str]:
    """返回 (有效日期, 创建时间, 日期来源说明)。"""
    ctime = datetime.fromtimestamp(path.stat().st_ctime)
    parsed_dates = parse_dates_from_filename(path.name)
    if parsed_dates:
        return max(parsed_dates), ctime, "文件名日期"
    return ctime.date(), ctime, "创建时间"


def list_newdata_excel_files(folder: Path) -> list[Path]:
    excel_files = list(folder.glob("*.xlsx")) + list(folder.glob("*.xls"))
    return sorted(
        [f for f in excel_files if not f.name.startswith("~$")],
        key=lambda p: p.name,
    )


def _format_newdata_file_info(
    path: Path,
    effective_date: date,
    ctime: datetime,
    source_label: str,
) -> str:
    return (
        f"{path.name}  （{source_label}：{effective_date.isoformat()}，"
        f"创建：{ctime.strftime('%Y-%m-%d %H:%M')}）"
    )


def select_newdata_file(folder: Path, *, non_interactive: bool = False) -> Path:
    """从 Newdata 文件夹中选取有效日期最新的单个 Excel 文件。"""
    candidates = list_newdata_excel_files(folder)
    if not candidates:
        raise FileNotFoundError(f"Newdata 文件夹中未找到任何 Excel 文件：{folder}")

    ranked: list[tuple[Path, date, datetime, str]] = []
    for path in candidates:
        effective_date, ctime, source_label = get_newdata_effective_date(path)
        ranked.append((path, effective_date, ctime, source_label))

    max_date = max(item[1] for item in ranked)
    latest = [item for item in ranked if item[1] == max_date]

    if len(latest) == 1:
        path, effective_date, ctime, source_label = latest[0]
        print(f"  使用新数据文件：{_format_newdata_file_info(path, effective_date, ctime, source_label)}")
        return path

    if non_interactive:
        path, effective_date, ctime, source_label = max(latest, key=lambda item: item[2])
        print(
            f"\n  非交互模式：{len(latest)} 个并列最新文件，"
            f"按创建时间自动选取：{_format_newdata_file_info(path, effective_date, ctime, source_label)}"
        )
        return path

    print(f"\n  发现 {len(latest)} 个并列最新的 Newdata 文件，请选择：")
    for idx, (path, effective_date, ctime, source_label) in enumerate(latest, start=1):
        print(f"    [{idx}] {_format_newdata_file_info(path, effective_date, ctime, source_label)}")

    while True:
        choice = input(f"  请输入序号 [1-{len(latest)}]: ").strip()
        if choice.isdigit():
            selected_idx = int(choice)
            if 1 <= selected_idx <= len(latest):
                path, effective_date, ctime, source_label = latest[selected_idx - 1]
                print(
                    f"  已选择新数据文件："
                    f"{_format_newdata_file_info(path, effective_date, ctime, source_label)}"
                )
                return path
        print("   请输入有效序号")


def read_excel_all_sheets(file_path: Path) -> pd.DataFrame:
    """读取单个 Excel 文件的全部 Sheet，合并为一个 DataFrame。"""
    all_data: list[pd.DataFrame] = []
    print(f"  正在读取新数据文件：{file_path.name}")
    try:
        xf = pd.ExcelFile(file_path)
        for sheet_name in xf.sheet_names:
            try:
                df = pd.read_excel(xf, sheet_name=sheet_name, dtype=str)
                if df.empty:
                    continue
                df.columns = [str(c).strip() for c in df.columns]
                for col in ["工单号", "用户编码", "所属馈线编码"]:
                    if col in df.columns:
                        df[col] = (
                            df[col].astype(str).str.strip().str.replace(".0", "", regex=False)
                        )
                df = df.dropna(how="all")
                all_data.append(df)
                print(f"    Sheet [{sheet_name}]：{len(df)} 行")
            except Exception as exc:
                print(f"    Sheet [{sheet_name}] 读取失败：{exc}", file=sys.stderr)
    except Exception as exc:
        raise RuntimeError(f"文件读取失败 [{file_path.name}]：{exc}") from exc
    finally:
        if "xf" in locals():
            xf.close()

    if not all_data:
        raise ValueError(f"Newdata 文件中未读取到任何有效数据：{file_path}")

    return pd.concat(all_data, ignore_index=True)


def read_folder_raw_data(folder_path: str, *, non_interactive: bool = False) -> pd.DataFrame:
    """从 Newdata 文件夹中选取最新 Excel 并读取其全部 Sheet。"""
    folder = Path(folder_path)
    if not folder.exists():
        raise FileNotFoundError(f"Newdata 文件夹不存在：{folder}")

    selected_file = select_newdata_file(folder, non_interactive=non_interactive)
    data = read_excel_all_sheets(selected_file)
    data.attrs["source_file"] = str(selected_file)
    return data


def _daily_report_set_complete(result_path: Path) -> bool:
    stem = result_path.stem.replace("停电用户_处理结果", "")
    required = (
        result_path.with_name(f"{stem}停电用户_处理结果_发出版.xlsx"),
        result_path.with_name(f"{stem}统计表.xlsx"),
        result_path.with_name(f"{stem}停电摘要（简版）.txt"),
        result_path.with_name(f"{stem}停电摘要（全量版）.txt"),
    )
    return all(path.is_file() and path.stat().st_size > 0 for path in required)


def _find_valid_result_file(search_dir: Path, *, require_complete: bool = False) -> Path | None:
    """在目录中查找最新且有效的 *_处理结果.xlsx（排除发出版、空文件）。"""
    candidates = [
        f
        for f in search_dir.glob("*.xlsx")
        if "处理结果" in f.name
        and "发出版" not in f.name
        and not f.name.startswith("~$")
        and f.stat().st_size > 100
        and (not require_complete or _daily_report_set_complete(f))
    ]
    if not candidates:
        return None
    return max(candidates, key=lambda f: f.stat().st_mtime)


def _folder_end_date(folder: Path) -> date | None:
    """从 output 子目录名解析数据截止日，取文件名/目录名中的最大日期。"""
    dates = parse_dates_from_filename(folder.name)
    return max(dates) if dates else None


def _resolve_latest_output_subdir(folder_path: Path, *, require_complete: bool = False) -> Path:
    """在 output 根目录下查找含有效处理结果、且数据截止日最新的子目录。"""
    subdirs = [p for p in folder_path.iterdir() if p.is_dir()]
    if not subdirs:
        return folder_path

    ranked: list[tuple[date, float, Path]] = []
    for subdir in subdirs:
        if _find_valid_result_file(subdir, require_complete=require_complete) is None:
            continue
        end_date = _folder_end_date(subdir)
        ranked.append(
            (
                end_date or date.min,
                subdir.stat().st_mtime,
                subdir,
            )
        )

    if ranked:
        chosen = max(ranked, key=lambda item: (item[0], item[1]))[2]
        print(f"  使用最新输出子目录：{chosen.name}")
        return chosen

    latest = max(subdirs, key=lambda p: p.stat().st_mtime)
    print(f"  使用最新输出子目录：{latest.name}")
    return latest


def read_existingdata_folder(folder_path: str) -> tuple[pd.DataFrame, Path | None]:
    """
    从 output/ 最新日期子目录中读取历史处理结果。
    - 自动查找 *_处理结果.xlsx（排除发出版），取最新修改的文件。
    - 读取其中"用户停电总次数统计表"Sheet，剥离分析附加列，还原为原始数据。
    返回：(历史原始数据 DataFrame, 处理结果文件路径（供累计预警使用）)
    """
    folder = Path(folder_path)
    if not folder.exists():
        print(f"  尚无历史处理结果目录，按首次日报运行：{folder}")
        return pd.DataFrame(), None

    search_dir = _resolve_latest_output_subdir(folder, require_complete=True)

    latest_file = _find_valid_result_file(search_dir, require_complete=True)
    if latest_file is None:
        print(
            f"  未找到历史处理结果文件，跳过历史窗口合并：{search_dir}",
            file=sys.stderr,
        )
        return pd.DataFrame(), None

    print(f"  使用历史处理结果文件：{latest_file.name}")

    try:
        df = pd.read_excel(latest_file, sheet_name="用户停电总次数统计表", dtype=str)
    except Exception as exc:
        print(
            f"  读取历史原始数据失败（文件可能已损坏，跳过）[{latest_file.name}]：{exc}",
            file=sys.stderr,
        )
        return pd.DataFrame(), None

    # 剥离脚本附加的分析列，还原为原始数据
    drop_cols = [c for c in RAW_ADDED_COLUMNS + ["本次较上次变化说明", "_user_feeder_key", "_是否统计_bool", "_不统计原因"] if c in df.columns]
    df = df.drop(columns=drop_cols)
    df.columns = [str(c).strip() for c in df.columns]
    df = df.dropna(how="all")

    for col in ["工单号", "用户编码", "所属馈线编码"]:
        if col in df.columns:
            df[col] = df[col].astype(str).str.strip().str.replace(".0", "", regex=False)

    print(f"  历史原始数据：{len(df)} 行")
    return df, latest_file


def format_date_range_str(start: date, end: date) -> str:
    """将日期范围格式化为 '2026.1.1-6.4' 形式。"""
    start_str = f"{start.year}.{start.month}.{start.day}"
    if start.year == end.year:
        end_str = f"{end.month}.{end.day}"
    else:
        end_str = f"{end.year}.{end.month}.{end.day}"
    return f"{start_str}-{end_str}"


def get_date_range_from_raw(df: pd.DataFrame) -> tuple[date | None, date | None]:
    """从原始数据中提取停电开始时间的最小值和最大值。"""
    ts = _outage_start_timestamps(df).dropna()
    if ts.empty:
        return None, None
    return ts.min().date(), ts.max().date()


def _outage_start_timestamps(df: pd.DataFrame) -> pd.Series:
    """优先去重后停电开始时间，否则停电开始时间。"""
    if df is None or df.empty:
        return pd.Series(dtype="datetime64[ns]")
    for col_name in ["去重后停电开始时间", "停电开始时间"]:
        if col_name in df.columns:
            return pd.to_datetime(df[col_name], errors="coerce")
    return pd.Series(pd.NaT, index=df.index)


def merge_and_deduplicate_raw(new_df: pd.DataFrame, existing_df: pd.DataFrame) -> pd.DataFrame:
    """按 Newdata 停电开始日闭区间整段替换历史；区间外历史保留，缺天不回填。"""
    if existing_df.empty:
        return new_df.copy()
    if new_df.empty:
        return existing_df.copy()

    new_ts = _outage_start_timestamps(new_df)
    valid = new_ts.dropna()
    if valid.empty:
        raise ValueError("Newdata 中无法解析停电开始时间，无法按日期窗口替换")

    window_start = valid.min().date()
    window_end = valid.max().date()
    exist_ts = _outage_start_timestamps(existing_df)
    exist_day = exist_ts.dt.date
    keep = exist_ts.isna() | (exist_day < window_start) | (exist_day > window_end)
    kept = existing_df.loc[keep].copy()
    dropped = int((~keep).sum())
    combined = pd.concat([kept, new_df], ignore_index=True)
    print(
        f"  Newdata 日期窗口：{window_start.isoformat()} ~ {window_end.isoformat()}"
        "（缺天按 0 条，不回填历史）"
    )
    print(
        f"  窗口外历史保留：{len(kept)} 行；窗口内历史丢弃：{dropped} 行；"
        f"Newdata：{len(new_df)} 行"
    )
    return combined.reset_index(drop=True)


def main() -> int:
    args = parse_args()
    runner = StepRunner(non_interactive=args.non_interactive)

    # ============================================================
    # 运行参数解析（命令行参数 > 顶部配置 > 默认值）
    # ============================================================
    period_start = parse_one_date(args.period_start) if args.period_start else None
    period_end = parse_one_date(args.period_end) if args.period_end else None
    if bool(args.period_start) != bool(args.period_end):
        print("--period-start 与 --period-end 必须同时指定", file=sys.stderr)
        return 1
    if period_start and period_end and period_end < period_start:
        print("--period-end 不得早于 --period-start", file=sys.stderr)
        return 1
    historical_mode = period_start is not None and period_end is not None

    today_explicitly_set = bool(args.today or TODAY_DATE) or historical_mode
    if historical_mode:
        today = period_end
    elif today_explicitly_set:
        today = parse_one_date(args.today or TODAY_DATE)
    else:
        today = date.today()
    if today is None:
        print("--today 日期格式应为 YYYY-MM-DD", file=sys.stderr)
        return 1

    major_dates_str = args.major_event_dates or MAJOR_EVENT_DATES_STR
    major_event_file = args.major_event_file or (MAJOR_EVENT_FILE_PATH or None)
    exclude_2025_dir = args.exclude_2025_dir or EXCLUDE_2025_DIR
    exclude_2026_dir = args.exclude_2026_dir or EXCLUDE_2026_DIR
    exclude_2025_file = args.exclude_2025_file
    exclude_2026_file = getattr(args, "exclude_2026_file", None)

    major_dates_2025, major_dates_2026, exclude_sources = load_exclude_major_dates(
        exclude_2025_dir,
        exclude_2026_dir,
        exclude_2025_file=exclude_2025_file,
        exclude_2026_file=exclude_2026_file,
    )
    extra_2026 = parse_date_list(major_dates_str) | read_major_event_dates(major_event_file)
    if extra_2026:
        major_dates_2026 |= extra_2026

    nea_ledger_file = args.nea_ledger or (NEA_LEDGER_PATH if Path(NEA_LEDGER_PATH).exists() else None)
    nea_gis_ids = read_nea_ledger(nea_ledger_file)

    newdata_dir      = args.newdata      or NEWDATA_FOLDER
    existingdata_dir = args.existingdata or EXISTINGDATA_FOLDER
    stats_tmpl       = STATS_TEMPLATE_PATH
    data_year = args.data_year
    annual_2025_path = Path(args.annual_2025_path) if args.annual_2025_path else Path(ANNUAL_2025_PATH)

    # 历史模式：按 data-year 解析输入文件
    resolved_input = args.input
    if historical_mode and data_year == 2025 and not resolved_input:
        resolved_input = str(annual_2025_path)

    print("\n" + "=" * 62)
    print("  供电可靠性停电统计分析脚本 v1.0.8")
    print("=" * 62)
    print(f"  今日日期     ：{today}{'（显式指定/区间终点）' if today_explicitly_set else '（暂定系统日期，将从数据自动推断）'}")
    if historical_mode:
        print(f"  历史区间     ：{period_start.isoformat()} ~ {period_end.isoformat()}")
        print(f"  主数据年份   ：{data_year or '（未指定，按输入推断）'}")
    print(
        f"  2025 剔除文件 ：{Path(exclude_sources['2025_file']).name}"
        f"（{len(major_dates_2025)} 天）"
    )
    print(
        f"  2026 剔除文件 ：{Path(exclude_sources['2026_file']).name}"
        f"（{len(major_dates_2026)} 天）"
    )
    if extra_2026:
        print(f"  2026 额外剔除 ：{len(extra_2026)} 天（命令行/配置补充）")
    print(f"  NEA 台账线路 ：{len(nea_gis_ids)} 条")
    print(f"  新数据目录   ：{newdata_dir}")
    print(f"  历史数据目录 ：{existingdata_dir}")
    print(f"  统计表模板   ：{stats_tmpl}")

    # ============================================================
    # 步骤 1：读取新数据
    # ============================================================
    raw_new: pd.DataFrame | None = None
    postgres_mode = (
        os.environ.get("DATA_BACKEND", "excel").strip().lower() == "postgres"
        and not historical_mode
        and not resolved_input
    )
    if postgres_mode:
        def _load_postgres():
            df = load_active_records()
            print(f"  PostgreSQL 当前有效数据：{len(df)} 行")
            return df
        raw_new = runner.run("读取 PostgreSQL 当前有效数据", _load_postgres)
    elif resolved_input:
        # 单文件模式（含 2025 Annual Summary）
        def _load_single():
            input_path = Path(resolved_input)
            if not input_path.exists():
                raise FileNotFoundError(f"输入文件不存在：{input_path}")
            if data_year == 2025 or "2025年停电用户数据" in input_path.name:
                df = read_2025_raw(str(input_path))
                df = normalize_2025_columns(df)
                print(f"  已读取 2025 年数据：{input_path.name}（{len(df)} 行）")
            else:
                df = pd.read_excel(input_path, sheet_name=normalize_sheet_arg(args.sheet))
                print(f"  已读取单文件：{input_path.name}（{len(df)} 行）")
            return df
        raw_new = runner.run("读取新数据（单文件模式）", _load_single)
    else:
        def _load_newdata():
            df = read_folder_raw_data(newdata_dir, non_interactive=args.non_interactive)
            print(f"  新数据合计：{len(df)} 行")
            return df
        raw_new = runner.run("读取新数据（Newdata 文件夹）", _load_newdata)

    if raw_new is None:
        print("[ERR]  新数据读取失败，程序终止", file=sys.stderr)
        return 1

    # ============================================================
    # 步骤 2：读取历史数据（历史区间模式跳过，避免混入区间外数据）
    # ============================================================
    existing_raw: pd.DataFrame = pd.DataFrame()
    previous_file_path: Path | None = None

    if not resolved_input and not historical_mode and not postgres_mode:
        def _load_existing():
            nonlocal previous_file_path
            df, pf = read_existingdata_folder(existingdata_dir)
            previous_file_path = pf
            return df
        result = runner.run("读取历史处理数据（output 最新子目录）", _load_existing, allow_skip=True)
        if result is not None:
            existing_raw = result
    elif postgres_mode:
        try:
            latest_dir = _resolve_latest_output_subdir(Path(existingdata_dir), require_complete=True)
            candidates = [
                path for path in latest_dir.glob("*处理结果.xlsx")
                if "_发出版" not in path.name and path.is_file() and _daily_report_set_complete(path)
            ]
            previous_file_path = max(candidates, key=lambda path: path.stat().st_mtime) if candidates else None
        except (FileNotFoundError, OSError):
            previous_file_path = None

    # ============================================================
    # 步骤 3：按 Newdata 日期窗口整段替换
    # ============================================================
    def _merge():
        df = merge_and_deduplicate_raw(raw_new, existing_raw)
        print(f"  合并后共 {len(df)} 行")
        return df

    raw = runner.run("按日期窗口替换合并原始数据", _merge)
    if raw is None:
        print("[ERR]  数据合并失败，程序终止", file=sys.stderr)
        return 1

    # ============================================================
    # 历史区间过滤
    # ============================================================
    if historical_mode:
        before = len(raw)
        raw = filter_raw_by_outage_date(raw, period_start, period_end)
        print(
            f"\n  区间过滤：{period_start.isoformat()} ~ {period_end.isoformat()}，"
            f"{before} → {len(raw)} 行"
        )
        if raw.empty:
            print("[ERR]  指定区间内无停电记录，程序终止", file=sys.stderr)
            return 1

    # ============================================================
    # 确定日期范围与输出路径
    # ============================================================
    if historical_mode:
        start_date, end_date = period_start, period_end
    else:
        start_date, end_date = get_date_range_from_raw(raw)

    # 若未显式指定 --today / TODAY_DATE，则自动使用数据最大停电日期作为"今日"
    if not today_explicitly_set and end_date is not None:
        today = end_date
        print(f"\n  今日日期（自动从数据推断）：{today}（数据截止日）")
    else:
        print(f"\n  今日日期：{today}")

    if start_date and end_date:
        date_range_str = format_date_range_str(start_date, end_date)
    else:
        date_range_str = f"{today.year}.{today.month}.{today.day}"
    print(f"  数据日期范围：{date_range_str}")

    if resolved_input and not historical_mode:
        # 单文件日常模式：沿用旧命名逻辑
        input_path = Path(resolved_input)
        output_path = Path(args.output) if args.output else input_path.with_name(f"{input_path.stem}_处理结果.xlsx")
        publish_output_path = Path(args.publish_output) if args.publish_output else default_publish_output_path(output_path)
        stats_output_path = Path(STATS_OUTPUT_PATH) if STATS_OUTPUT_PATH else output_path.with_name(f"{output_path.stem.replace('_处理结果', '')}统计表.xlsx")
    else:
        # 文件夹模式 / 历史区间：output/<日期范围>/
        out_base = Path(args.output).parent if args.output else Path(OUTPUT_BASE_DIR) / date_range_str
        if out_base.exists():
            if out_base.is_dir():
                pass  # 目录已存在，无需创建
            else:
                print(f"  ⚠ 输出路径存在但不是目录，将自动删除并重建：{out_base}")
                out_base.unlink()
                os.makedirs(out_base, exist_ok=True)
        else:
            os.makedirs(out_base, exist_ok=True)
        output_path       = Path(args.output)        if args.output        else out_base / f"{date_range_str}停电用户_处理结果.xlsx"
        publish_output_path = Path(args.publish_output) if args.publish_output else out_base / f"{date_range_str}停电用户_处理结果_发出版.xlsx"
        stats_output_path = Path(STATS_OUTPUT_PATH)  if STATS_OUTPUT_PATH  else out_base / f"{date_range_str}统计表.xlsx"

    # A correction with the same date range must not overwrite the prior
    # successful workbook: it is the comparison and rollback baseline.
    if not resolved_input and not historical_mode and output_path.exists():
        token = datetime.now().strftime("%Y%m%d%H%M%S%f")
        output_path = output_path.with_name(f"{date_range_str}_{token}停电用户_处理结果.xlsx")
        publish_output_path = publish_output_path.with_name(
            f"{date_range_str}_{token}停电用户_处理结果_发出版.xlsx"
        )
        stats_output_path = stats_output_path.with_name(f"{date_range_str}_{token}统计表.xlsx")
        print(f"  同日期范围再次运行，使用版本化文件名：{token}")
    print(f"  处理结果输出  ：{output_path}")
    print(f"  发出版输出    ：{publish_output_path}")
    print(f"  统计表输出    ：{stats_output_path}")
    summary_stem = output_path.stem.replace("停电用户_处理结果", "")
    summary_simple_path = output_path.parent / f"{summary_stem}停电摘要（简版）.txt"
    summary_full_path = output_path.parent / f"{summary_stem}停电摘要（全量版）.txt"

    # 历史预警记录
    if args.previous_output:
        previous_repeated_lines = read_previous_repeated_lines(args.previous_output)
    elif previous_file_path and not historical_mode:
        previous_repeated_lines = read_previous_repeated_lines(str(previous_file_path))
    else:
        previous_repeated_lines = pd.DataFrame(columns=REPEATED_LINE_COLUMNS)

    # 主分析使用的重大事件日：按 data-year 选择
    if data_year == 2025:
        major_dates_main = major_dates_2025
    else:
        major_dates_main = major_dates_2026

    # ============================================================
    # 步骤 4：可靠性规则分析
    # ============================================================
    def _analyze():
        tbls, sts = build_tables(raw, today, major_dates_main, previous_repeated_lines, nea_gis_ids)
        print(f"  统计用户数：{sts.get('统计用户数', 0)}")
        print(f"  频繁停电用户：{sts.get('频繁停电用户数', 0)}，线路：{sts.get('频繁停电线路数', 0)}")
        print(f"  停电预警用户：{sts.get('停电预警用户数', 0)}，线路：{sts.get('停电预警线路数', 0)}")
        return tbls, sts

    result_analyze = runner.run("可靠性规则分析", _analyze)
    if result_analyze is None:
        print("[ERR]  分析失败，程序终止", file=sys.stderr)
        return 1
    tables, stats = result_analyze
    current_stats_dicts = compute_2026_stats_dicts(tables)
    comparison, change_detail, user_explanations = compare_reports(
        tables, previous_file_path, tables["用户停电总次数统计表"], existing_raw,
        data_year=today.year,
        current_source=str(raw_new.attrs.get("source_file", resolved_input or newdata_dir)),
        current_stats_dicts=current_stats_dicts,
    )
    raw_detail = tables["用户停电总次数统计表"].copy()
    raw_detail["本次较上次变化说明"] = raw_detail["用户编码"].map(
        lambda value: user_explanations.get(normalize_identifier(value), "")
    )
    tables["用户停电总次数统计表"] = raw_detail
    tables["变化明细"] = change_detail
    comparison_path = output_path.with_name(f"{output_path.stem}_比较.json")
    comparison["change_detail_sheet"] = "变化明细"
    comparison["current_report"] = str(output_path)
    write_text_atomic(comparison_path, json.dumps(comparison, ensure_ascii=False, indent=2))
    if comparison.get("available"):
        print(f"  较上次日报变化：{comparison['metrics']}")
    else:
        print(f"  较上次日报变化：{comparison.get('reason', '暂无可比基准')}")

    # ============================================================
    # 步骤 5：输出处理结果 Excel
    # ============================================================
    def _write_result():
        write_tables(tables, output_path)
        print(f"  已写入：{output_path}")

    runner.run("写入停电用户_处理结果.xlsx", _write_result)

    # ============================================================
    # 步骤 6：输出发出版 Excel
    # ============================================================
    def _write_publish():
        write_tables(build_publish_tables(tables), publish_output_path)
        print(f"  已写入：{publish_output_path}")

    runner.run("写入停电用户_处理结果_发出版.xlsx", _write_publish)

    # ============================================================
    # 步骤 7：生成停电摘要文字
    # ============================================================
    def _gen_summary():
        stats_dicts = compute_2026_stats_dicts(tables)
        s_simple = build_summary_text(
            tables,
            today,
            stats_dicts,
            data_end_date=end_date,
            all_show_line_names=False,
        )
        s_full = build_summary_text(
            tables,
            today,
            stats_dicts,
            data_end_date=end_date,
            all_show_line_names=True,
        )
        if comparison.get("available"):
            comparison_line = format_comparison_summary(comparison)
            s_simple = s_simple.rstrip("\n") + "\n" + comparison_line + "\n"
            s_full = s_full.rstrip("\n") + "\n" + comparison_line + "\n"
        print("\n" + summary_banner("停电摘要（简版）"))
        print(s_simple, end="" if s_simple.endswith("\n") else "\n")
        print("\n" + summary_banner("停电摘要（全量版）"))
        print(s_full, end="" if s_full.endswith("\n") else "\n")
        write_text_atomic(summary_simple_path, s_simple)
        write_text_atomic(summary_full_path, s_full)
        print(f"  简版摘要输出  ：{summary_simple_path}")
        print(f"  全量摘要输出  ：{summary_full_path}")
        return s_simple, s_full

    runner.run("生成停电摘要文字", _gen_summary, allow_skip=True)

    # ============================================================
    # 步骤 8：生成统计表
    # ============================================================
    def _gen_stats():
        if not Path(stats_tmpl).exists():
            raise FileNotFoundError(f"统计表模板不存在：{stats_tmpl}")
        raw_2025_for_stats = str(annual_2025_path) if annual_2025_path.exists() else RAW_2025_DATA_PATH
        generate_statistics_table(
            tables_2026=tables,
            template_path=stats_tmpl,
            output_path=stats_output_path,
            raw_2025_path=raw_2025_for_stats,
            period_start=start_date,
            period_end=end_date,
            major_dates_2026=major_dates_2026,
            major_dates_2025=major_dates_2025,
            nea_gis_ids=nea_gis_ids,
        )

    runner.run("生成统计表", _gen_stats, allow_skip=True)

    # ============================================================
    # 验证摘要
    # ============================================================
    if args.validate:
        print("\n── 处理摘要 ──")
        for key, value in stats.items():
            print(f"  {key}: {value}")
        print(f"  今日日期: {today.isoformat()}")
        print(f"  2025 剔除文件: {exclude_sources['2025_file']}")
        print(f"  2025 重大事件停电日数量: {len(major_dates_2025)}")
        print(f"  2026 剔除文件: {exclude_sources['2026_file']}")
        print(f"  2026 重大事件停电日数量: {len(major_dates_2026)}")
        if extra_2026:
            print(f"  2026 额外剔除日数量: {len(extra_2026)}")

    runner.summary()
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
